"""
ProfitPulse Brief Builder
Target: Affinity MSP | Date: 27 Aug 2026
Three slide prospect facing deck. House style per Section 6. Brand colours only. Zero dashes.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand colours, the seven only
BLACK     = RGBColor(0x00, 0x00, 0x00)
TEAL      = RGBColor(0x01, 0xA2, 0x96)
AMBER_B   = RGBColor(0xF8, 0xC8, 0x06)
AMBER_D   = RGBColor(0xF6, 0xA1, 0x02)
GOLD      = RGBColor(0xE3, 0xA7, 0x12)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xE6, 0xE5, 0xDE)

SERIF = "Cambria"
SANS  = "Calibri"

W = Inches(13.333)
H = Inches(7.5)

DATE_STR = "27 Aug 2026"
COMPANY = "Affinity MSP"
OUT_PATH = "/home/user/Daily-Outreach/Out-reach efforts/Brief_AffinityMSP_27Aug2026.pptx"


def bg(slide, colour):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = colour


def rect(slide, left, top, width, height, colour, line=None, line_w=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = colour
    if line:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    # Strip the theme style reference entirely, it carries a preset drop shadow
    # (effectRef) that LibreOffice applies even when spPr has an empty effectLst.
    style_el = shp._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style_el is not None:
        shp._element.remove(style_el)
    return shp


def text(slide, s, left, top, width, height, size=12, bold=False, colour=WHITE,
         align=PP_ALIGN.LEFT, font=SANS, italic=False, anchor=None, line_spacing=None,
         shrink=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if shrink:
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.NONE
    if anchor:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = s
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = colour
    return box


def multiline(slide, lines, left, top, width, height, font=SANS, size=12,
              colour=OFF_WHITE, align=PP_ALIGN.LEFT, space_after=4, line_spacing=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for ln in lines:
        cfg = {"text": ln, "size": size, "colour": colour, "bold": False, "italic": False} \
            if isinstance(ln, str) else {
                "text": ln.get("text", ""), "size": ln.get("size", size),
                "colour": ln.get("colour", colour), "bold": ln.get("bold", False),
                "italic": ln.get("italic", False),
            }
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = cfg["text"]
        r.font.name = font
        r.font.size = Pt(cfg["size"])
        r.font.bold = cfg["bold"]
        r.font.italic = cfg["italic"]
        r.font.color.rgb = cfg["colour"]
    return box


def invisible_link(slide, left, top, width, height, url):
    """An invisible clickable overlay, so linked text keeps its authored brand colour
    instead of LibreOffice/PowerPoint's theme hyperlink colour."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.background()
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.click_action.hyperlink.address = url
    return shp


def chrome(slide, eyebrow, right_label=BLACK):
    """Left accent stripe, black header band, footer line. Fixed chrome per 6.0A."""
    bg(slide, WHITE)
    # Left accent stripe, amber, full height
    rect(slide, Inches(0), Inches(0), Inches(0.1), H, AMBER_B)
    # Header band, black, 1 inch tall
    rect(slide, Inches(0.1), Inches(0), W - Inches(0.1), Inches(1.0), BLACK)
    text(slide, eyebrow, Inches(0.35), Inches(0.32), Inches(8.5), Inches(0.4),
         size=12, bold=True, colour=OFF_WHITE, font=SANS)
    text(slide, "PROFITPULSE", Inches(9.0), Inches(0.32), Inches(4.0), Inches(0.4),
         size=13, bold=True, colour=TEAL, align=PP_ALIGN.RIGHT, font=SANS)
    # Footer line
    rect(slide, Inches(0.35), Inches(7.05), Inches(12.6), Pt(0.75), TEAL)
    text(slide, "Prepared by Nitesh Roopa CA, Managing Partner and Founder, ProfitPulse, Profit-Pulse.com.au",
         Inches(0.35), Inches(7.1), Inches(9.5), Inches(0.3), size=8, colour=BLACK)
    text(slide, DATE_STR, Inches(10.5), Inches(7.1), Inches(2.45), Inches(0.3),
         size=8, colour=BLACK, align=PP_ALIGN.RIGHT)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 1 — THE COMMERCIAL INTELLIGENCE BRIEF
# ═══════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank)
chrome(s1, "COMMERCIAL INTELLIGENCE BRIEF")

text(s1, "Affinity MSP", Inches(0.35), Inches(1.18), Inches(9.5), Inches(0.75),
     size=38, bold=True, colour=BLACK, font=SERIF)
text(s1, "Managed IT services provider, Mount Waverley, Melbourne VIC",
     Inches(0.35), Inches(1.85), Inches(11.5), Inches(0.35), size=13, colour=BLACK)

# Stat cards row, 6 cards, even spacing, top y=2.4, height 1.6
cards = [
    ("$6.6M", "FY revenue", "Smart50 2025 citation"),
    ("42%", "Three year revenue growth", "Smart50 2025 citation"),
    ("36", "Team size, employees", "Smart50 2025 citation"),
    ("#29", "Smart50 2025 national rank", "SmartCompany Nov 2025"),
    ("#1 AU", "MSP 501 Australia rank", "Channel Futures 2026"),
    ("2019", "Year founded", "Smart50 2025 citation"),
]
n = len(cards)
gap = Inches(0.15)
total_gap = gap * (n - 1)
usable = W - Inches(0.35) - Inches(0.35)
card_w = Emu(int((usable - total_gap) / n))
card_top = Inches(2.4)
card_h = Inches(1.6)
x = Inches(0.35)
for num, label, src in cards:
    rect(s1, x, card_top, card_w, card_h, BLACK)
    rect(s1, x, card_top, card_w, Pt(3), TEAL)
    text(s1, num, x + Inches(0.12), card_top + Inches(0.18), card_w - Inches(0.24), Inches(0.55),
         size=27, bold=True, colour=AMBER_B, font=SERIF)
    text(s1, label, x + Inches(0.12), card_top + Inches(0.78), card_w - Inches(0.24), Inches(0.5),
         size=10, colour=WHITE, line_spacing=1.0)
    text(s1, src, x + Inches(0.12), card_top + Inches(1.32), card_w - Inches(0.24), Inches(0.25),
         size=7, colour=OFF_WHITE, italic=True)
    x = x + card_w + gap

# Section label
text(s1, "KEY COMMERCIAL SIGNALS", Inches(0.35), Inches(4.25), Inches(6), Inches(0.3),
     size=12, bold=True, colour=TEAL)

signals = [
    "Ranked 29th, Smart50 2025, revenue $6.6M, 42% three year growth (SmartCompany, Nov 2025)",
    "Named #1 Australian MSP, #69 globally, 2026 Channel Futures MSP 501 (Channel Futures, 2026)",
    "36 staff on a utilisation based delivery model (Smart50 2025 award citation)",
    "Built without external investors, has avoided pursuing mergers (Smart50 2025 coverage)",
    "Founded 2019, Mount Waverley VIC, about 17km from Melbourne CBD (company site, directories)",
]
multiline(s1, signals, Inches(0.35), Inches(4.6), Inches(12.6), Inches(2.35),
          size=12, colour=BLACK, space_after=10, line_spacing=1.05)

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 2 — THE OPPORTUNITY
# ═══════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank)
bg(s2, WHITE)
rect(s2, Inches(0), Inches(0), Inches(0.1), H, AMBER_B)
rect(s2, Inches(0.1), Inches(0), W - Inches(0.1), Inches(1.0), BLACK)
text(s2, "THE OPPORTUNITY", Inches(0.35), Inches(0.2), Inches(8), Inches(0.35),
     size=12, bold=True, colour=OFF_WHITE)
text(s2, "Affinity MSP: three commercial observations from ProfitPulse",
     Inches(0.35), Inches(0.53), Inches(9.5), Inches(0.35), size=11, colour=TEAL)
text(s2, "PROFITPULSE", Inches(9.0), Inches(0.32), Inches(4.0), Inches(0.4),
     size=13, bold=True, colour=TEAL, align=PP_ALIGN.RIGHT)
rect(s2, Inches(0.35), Inches(7.05), Inches(12.6), Pt(0.75), TEAL)
text(s2, "Prepared by Nitesh Roopa CA, Managing Partner and Founder, ProfitPulse, Profit-Pulse.com.au",
     Inches(0.35), Inches(7.1), Inches(9.5), Inches(0.3), size=8, colour=BLACK)
text(s2, DATE_STR, Inches(10.5), Inches(7.1), Inches(2.45), Inches(0.3),
     size=8, colour=BLACK, align=PP_ALIGN.RIGHT)

cols = [
    dict(fill=TEAL, txt=BLACK, idx_c=BLACK,
         idx="01", head="Growth is outrunning utilisation visibility",
         body=("Affinity MSP has grown headcount to 36 while lifting revenue 42 percent "
               "over three years, and has just been named the top ranked Australian "
               "provider on the 2026 MSP 501 index. In a utilisation based model, that "
               "pace of hiring makes it easy for billable mix to drift lower before "
               "anyone notices.")),
    dict(fill=BLACK, txt=WHITE, idx_c=AMBER_B,
         idx="02", head="A bootstrapped model concentrates financial risk",
         body=("The business has grown without external investors and has reportedly "
               "chosen not to pursue mergers, protecting its independence. That "
               "discipline is admirable, but every capital and pricing decision runs "
               "through a founder led team with no dedicated finance function testing "
               "the numbers independently.")),
    dict(fill=GOLD, txt=BLACK, idx_c=BLACK,
         idx="03", head="External recognition should now trigger internal rigour",
         body=("Two award programs in twelve months, Smart50 nationally and MSP 501 "
               "globally, have validated growth and operational depth to outsiders. "
               "The next test is whether revenue per technician, capacity headroom, "
               "and margin by client are tracked with the same rigour the judges "
               "applied.")),
]
col_gap = Inches(0.12)
col_w = Emu(int((W - Inches(0.35) - Inches(0.35) - col_gap * 2) / 3))
col_top = Inches(1.2)
col_h = Inches(4.85)
x = Inches(0.35)
for c in cols:
    rect(s2, x, col_top, col_w, col_h, c["fill"])
    text(s2, c["idx"], x + Inches(0.25), col_top + Inches(0.2), col_w - Inches(0.5), Inches(0.7),
         size=30, bold=True, colour=c["idx_c"], font=SERIF)
    text(s2, c["head"], x + Inches(0.25), col_top + Inches(1.0), col_w - Inches(0.5), Inches(0.9),
         size=15, bold=True, colour=c["txt"], font=SERIF, line_spacing=1.05)
    text(s2, c["body"], x + Inches(0.25), col_top + Inches(1.95), col_w - Inches(0.5), Inches(2.75),
         size=10.5, colour=c["txt"], line_spacing=1.15)
    x = x + col_w + col_gap

text(s2, "These are observations offered in good faith. The business has built something "
         "impressive. The question is simply whether the financial architecture matches the ambition.",
     Inches(0.35), Inches(6.15), Inches(12.6), Inches(0.75), size=11, italic=True,
     colour=BLACK, line_spacing=1.1)

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 3 — THE RECOMMENDATION
# ═══════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank)
chrome(s3, "THE RECOMMENDATION")

left_x = Inches(0.35)
left_w = Inches(7.6)

text(s3, "Workforce Capacity and Utilisation Review",
     left_x, Inches(1.15), left_w, Inches(0.95), size=21, bold=True, colour=BLACK,
     font=SERIF, line_spacing=1.05)
text(s3, "$3,950 one off", left_x, Inches(2.18), left_w, Inches(0.4),
     size=16, bold=True, colour=TEAL)
text(s3,
     "A four week review of billable mix, revenue per technician, and capacity "
     "headroom, so the next hire lands exactly where growth needs it.",
     left_x, Inches(2.62), left_w, Inches(0.65), size=12, colour=BLACK,
     line_spacing=1.15)

# Step one block
rect(s3, left_x, Inches(3.4), left_w, Inches(1.3), OFF_WHITE)
text(s3, "Step one, answer a few quick questions", left_x + Inches(0.2), Inches(3.55),
     left_w - Inches(0.4), Inches(0.35), size=13, bold=True, colour=BLACK)
text(s3, "See the solutions matched to your size and industry.",
     left_x + Inches(0.2), Inches(3.92), left_w - Inches(0.4), Inches(0.35),
     size=11, colour=BLACK)
q_box = text(s3, "profit-pulse.com.au/services/find-your-fit", left_x + Inches(0.2), Inches(4.28),
     left_w - Inches(0.4), Inches(0.35), size=12, bold=True, colour=TEAL)
invisible_link(s3, left_x + Inches(0.2), Inches(4.28), Inches(4.2), Inches(0.32),
    "https://profit-pulse.com.au/services/find-your-fit/?utm_source=outreach&utm_medium=pptx&utm_campaign=nightly_outreach&utm_content=affinity_msp")

# Direct CTA button
btn = rect(s3, left_x, Inches(4.95), Inches(6.2), Inches(0.5), AMBER_D)
btn_tf = btn.text_frame
btn_tf.word_wrap = False
btn_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
btn_tf.margin_left = Pt(2)
btn_tf.margin_right = Pt(2)
bp = btn_tf.paragraphs[0]
bp.alignment = PP_ALIGN.CENTER
br = bp.add_run()
br.text = "Purchase the suggested product now to get started"
br.font.name = SANS
br.font.size = Pt(12)
br.font.bold = True
br.font.color.rgb = BLACK
btn.click_action.hyperlink.address = "https://buy.stripe.com/eVq9AU2uA66cbdu9wr3ks1F"

text(s3, "Prefer a conversation first?", left_x, Inches(5.7), left_w, Inches(0.3),
     size=11, colour=BLACK)
call_box = text(s3, "Book a complimentary discovery call", left_x, Inches(6.02), left_w, Inches(0.35),
     size=12, bold=True, colour=TEAL)
invisible_link(s3, left_x, Inches(6.02), Inches(2.8), Inches(0.32),
    "https://bookings.cloud.microsoft/book/ProfitPulse1@profit-pulse.com.au/?ismsaljsauthenabled=true")

# Right column, credibility panel
rx = Inches(8.15)
rw = Inches(4.85)
rect(s3, rx, Inches(1.2), rw, Inches(5.35), BLACK)
rect(s3, rx, Inches(1.2), rw, Pt(3), TEAL)
text(s3, "Nitesh Roopa", rx + Inches(0.25), Inches(1.45), rw - Inches(0.5), Inches(0.4),
     size=17, bold=True, colour=AMBER_B, font=SERIF)
text(s3, "CA, Managing Partner, ProfitPulse", rx + Inches(0.25), Inches(1.85), rw - Inches(0.5), Inches(0.35),
     size=11.5, colour=WHITE)
rect(s3, rx + Inches(0.25), Inches(2.3), rw - Inches(0.5), Pt(1.25), TEAL)

cred = [
    "16 years across 4 countries",
    "52 deals executed and managed",
    "Largest single deal, USD 1.3 billion, Cahora Bassa",
    "Total GRBT project value over AUD 10 billion",
]
multiline(s3, cred, rx + Inches(0.25), Inches(2.5), rw - Inches(0.5), Inches(1.9),
          size=10.5, colour=OFF_WHITE, space_after=8, line_spacing=1.1)

contact = [
    {"text": "Profit-Pulse.com.au", "colour": OFF_WHITE},
    {"text": "Nitesh@Profit-Pulse.com.au", "colour": TEAL},
    {"text": "+61 411 876 267", "colour": OFF_WHITE},
    {"text": "linkedin.com/in/nitesh-roopa-77594163", "colour": TEAL},
]
multiline(s3, contact, rx + Inches(0.25), Inches(4.55), rw - Inches(0.5), Inches(1.8),
          size=10.5, space_after=7, line_spacing=1.1)

prs.save(OUT_PATH)
print("PPTX saved:", OUT_PATH)

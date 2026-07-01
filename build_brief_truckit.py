"""
ProfitPulse Brief Builder
Target: Truckit.net | Date: 02 Jul 2026
Three slide prospect facing deck. House style per Section 6.0/6.0A.
White background, black header band, amber left accent stripe, stat cards.
Brand colours only. Zero dashes.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand colours
BLACK      = RGBColor(0x00, 0x00, 0x00)
TEAL       = RGBColor(0x01, 0xA2, 0x96)
AMBER_B    = RGBColor(0xF8, 0xC8, 0x06)
AMBER_D    = RGBColor(0xF6, 0xA1, 0x02)
GOLD       = RGBColor(0xE3, 0xA7, 0x12)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE  = RGBColor(0xE6, 0xE5, 0xDE)

W = Inches(13.333)
H = Inches(7.5)
LEFT_MARGIN = Inches(0.62)
RIGHT_EDGE = Inches(12.72)
CONTENT_W = RIGHT_EDGE - LEFT_MARGIN

FONT_HEAD = "Georgia"
FONT_BODY = "Calibri"

DATE_STAMP = "02 Jul 2026"
COMPANY = "Truckit.net"


def set_background(slide, colour):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = colour


def _strip_style(shape):
    """Remove the p:style element (theme lnRef/fillRef/effectRef/fontRef) so
    LibreOffice cannot pull in a default shadow or off brand accent colour."""
    el = shape._element
    style = el.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        el.remove(style)


def add_rect(slide, left, top, width, height, fill_colour, line_colour=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_colour
    if line_colour:
        shape.line.color.rgb = line_colour
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    _strip_style(shape)
    return shape


def add_text(slide, text, left, top, width, height, font_name=FONT_BODY,
             font_size=12, bold=False, colour=BLACK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, line_spacing=1.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if line_spacing != 1.0:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = colour
    return box


def add_hyperlink_text(slide, text, url, left, top, width, height,
                        font_name=FONT_BODY, font_size=12, bold=False,
                        colour=TEAL, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = colour
    run.font.underline = False
    run.hyperlink.address = url
    return box


def add_chrome(slide, eyebrow, right_label=COMPANY.upper()):
    """Fixed chrome: left accent stripe, header band, footer line."""
    set_background(slide, WHITE)
    # Left accent stripe
    add_rect(slide, Inches(0), Inches(0), Inches(0.1), H, AMBER_D)
    # Header band
    add_rect(slide, Inches(0), Inches(0), W, Inches(1.0), BLACK)
    add_text(slide, eyebrow, LEFT_MARGIN, Inches(0.36), Inches(7.5), Inches(0.35),
              font_name=FONT_BODY, font_size=12, bold=True, colour=OFF_WHITE)
    add_text(slide, right_label, Inches(9.0), Inches(0.36), Inches(3.72), Inches(0.35),
              font_name=FONT_BODY, font_size=12, bold=True, colour=TEAL, align=PP_ALIGN.RIGHT)
    # Footer line
    add_rect(slide, LEFT_MARGIN, Inches(7.0), CONTENT_W, Pt(0.75), RGBColor(0xCC, 0xCC, 0xCC))
    add_text(slide, "Prepared by Nitesh Roopa CA, Managing Partner and Founder, ProfitPulse, Profit-Pulse.com.au",
              LEFT_MARGIN, Inches(7.08), Inches(9.5), Inches(0.3),
              font_name=FONT_BODY, font_size=8, colour=RGBColor(0x77, 0x77, 0x77))
    add_text(slide, DATE_STAMP, Inches(10.5), Inches(7.08), Inches(2.22), Inches(0.3),
              font_name=FONT_BODY, font_size=8, colour=RGBColor(0x77, 0x77, 0x77), align=PP_ALIGN.RIGHT)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]

# ═══════════════════════════════════════════════════════════════════
# SLIDE 1: COMMERCIAL INTELLIGENCE BRIEF
# ═══════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank)
add_chrome(s1, "COMMERCIAL INTELLIGENCE BRIEF")

add_text(s1, COMPANY, LEFT_MARGIN, Inches(1.15), Inches(9), Inches(0.65),
          font_name=FONT_HEAD, font_size=40, bold=True, colour=BLACK)
add_text(s1, "Online freight marketplace, Milton, Brisbane QLD",
          LEFT_MARGIN, Inches(1.85), Inches(10), Inches(0.35),
          font_name=FONT_BODY, font_size=13, colour=RGBColor(0x33, 0x33, 0x33))

# Stat cards: 5 cards, 2.3in wide, 0.15in gap, top y=2.35, height 1.6
CARD_W = Inches(2.3)
CARD_GAP = Inches(0.15)
CARD_TOP = Inches(2.35)
CARD_H = Inches(1.6)
cards = [
    ("$31M", "Revenue\nFY2025", "SmartCompany Smart50 2025"),
    ("21%", "Revenue\ngrowth", "SmartCompany Smart50 2025"),
    ("13", "Team\nmembers", "SmartCompany Smart50 2025"),
    ("2013", "Company\nfounded", "Truckit.net company site"),
    ("2X", "Smart50\nhonouree", "SmartCompany 2023, 2025"),
]
x = LEFT_MARGIN
for number, label, source in cards:
    add_rect(s1, x, CARD_TOP, CARD_W, CARD_H, BLACK)
    add_rect(s1, x, CARD_TOP, CARD_W, Pt(4), TEAL)
    add_text(s1, number, x + Inches(0.15), CARD_TOP + Inches(0.18), CARD_W - Inches(0.3), Inches(0.5),
              font_name=FONT_HEAD, font_size=28, bold=True, colour=AMBER_B)
    add_text(s1, label, x + Inches(0.15), CARD_TOP + Inches(0.72), CARD_W - Inches(0.3), Inches(0.5),
              font_name=FONT_BODY, font_size=11, colour=OFF_WHITE, line_spacing=1.05)
    add_text(s1, source, x + Inches(0.15), CARD_TOP + Inches(1.28), CARD_W - Inches(0.3), Inches(0.28),
              font_name=FONT_BODY, font_size=7, colour=RGBColor(0x99, 0x99, 0x99))
    x += CARD_W + CARD_GAP

# Key Commercial Signals
add_text(s1, "KEY COMMERCIAL SIGNALS", LEFT_MARGIN, Inches(4.18), Inches(6), Inches(0.3),
          font_name=FONT_BODY, font_size=12, bold=True, colour=TEAL)

signals = [
    "Ranked 47th on SmartCompany Smart50 2025, its second appearance since 2023.",
    "Described by SmartCompany as Australia's largest freight marketplace.",
    "States plans to scale using AI and machine learning, first in Australia then globally.",
    "Runs with 13 people generating $31 million in revenue, a lean marketplace model.",
    "Ranked 28th on SmartCompany Smart50 2023, a two year track record of growth.",
    "Founded 2013 by Robert Russell after a decade working in UK logistics.",
]
y = Inches(4.55)
for sig in signals:
    add_rect(s1, LEFT_MARGIN, y + Inches(0.08), Inches(0.09), Inches(0.09), AMBER_D)
    add_text(s1, sig, LEFT_MARGIN + Inches(0.24), y, CONTENT_W - Inches(0.24), Inches(0.36),
              font_name=FONT_BODY, font_size=11, colour=BLACK)
    y += Inches(0.375)

# ═══════════════════════════════════════════════════════════════════
# SLIDE 2: THE OPPORTUNITY
# ═══════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank)
add_chrome(s2, "THE OPPORTUNITY")

add_text(s2, "Truckit.net: three commercial observations from ProfitPulse",
          LEFT_MARGIN, Inches(1.15), Inches(11), Inches(0.4),
          font_name=FONT_HEAD, font_size=17, bold=True, colour=BLACK)

COL_W = Inches(3.9)
COL_GAP = Inches(0.2)
COL_TOP = Inches(1.75)
COL_H = Inches(4.65)

observations = [
    (TEAL, BLACK, "01", "Thin team, big throughput",
     "Thirteen people are generating 31 million dollars in revenue, a throughput far above typical logistics businesses of this size. That kind of ratio usually means the business is a broker of freight volume rather than a heavy asset owner, which makes understanding the true margin behind that revenue the first question worth answering."),
    (BLACK, OFF_WHITE, "02", "Two time Smart50 honouree",
     "A rank of 28th in 2023 and 47th in 2025 confirms a genuine multi year growth story rather than a single good year. Consistent recognition like this tends to attract inbound deal and partnership interest, which raises the value of having board grade numbers ready before those conversations happen."),
    (GOLD, BLACK, "03", "A stated global ambition",
     "The plan to scale using machine learning and AI, first in Australia then globally, is a step change from steady growth. International expansion multiplies capital, cash flow and capacity questions overnight, and a costed twelve month growth plan turns that ambition into a funded, sequenced plan rather than a headline."),
]

x = LEFT_MARGIN
for fill, text_colour, idx, header, para in observations:
    add_rect(s2, x, COL_TOP, COL_W, COL_H, fill)
    add_text(s2, idx, x + Inches(0.22), COL_TOP + Inches(0.2), COL_W - Inches(0.44), Inches(0.6),
              font_name=FONT_HEAD, font_size=30, bold=True, colour=text_colour)
    add_text(s2, header, x + Inches(0.22), COL_TOP + Inches(0.95), COL_W - Inches(0.44), Inches(0.55),
              font_name=FONT_BODY, font_size=14, bold=True, colour=text_colour, line_spacing=1.05)
    add_text(s2, para, x + Inches(0.22), COL_TOP + Inches(1.65), COL_W - Inches(0.44), Inches(2.85),
              font_name=FONT_BODY, font_size=11, colour=text_colour, line_spacing=1.15)
    x += COL_W + COL_GAP

add_text(s2, "These observations are offered in good faith. Truckit.net has built something genuinely lean and fast. The question is simply whether the financial architecture keeps pace with the ambition.",
          LEFT_MARGIN, Inches(6.55), CONTENT_W, Inches(0.4),
          font_name=FONT_BODY, font_size=10, italic=True, colour=RGBColor(0x44, 0x44, 0x44))

# ═══════════════════════════════════════════════════════════════════
# SLIDE 3: THE RECOMMENDATION AND HOW TO START
# ═══════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank)
add_chrome(s3, "THE RECOMMENDATION")

LEFT_COL_W = Inches(7.3)
RIGHT_COL_X = LEFT_MARGIN + LEFT_COL_W + Inches(0.3)
RIGHT_COL_W = RIGHT_EDGE - RIGHT_COL_X

# Left column: recommendation
add_text(s3, "Strategic Growth Diagnostic", LEFT_MARGIN, Inches(1.15), LEFT_COL_W, Inches(0.5),
          font_name=FONT_HEAD, font_size=24, bold=True, colour=BLACK)
add_text(s3, "$5,000 one off", LEFT_MARGIN, Inches(1.68), LEFT_COL_W, Inches(0.4),
          font_name=FONT_BODY, font_size=16, bold=True, colour=AMBER_D)
add_text(s3, "Six weeks to map revenue, capacity and margin headroom into a funded, scenario tested twelve month growth plan.",
          LEFT_MARGIN, Inches(2.15), LEFT_COL_W, Inches(0.65),
          font_name=FONT_BODY, font_size=12, colour=BLACK, line_spacing=1.15)

# Step one block
add_rect(s3, LEFT_MARGIN, Inches(2.95), LEFT_COL_W, Inches(1.55), RGBColor(0xF2, 0xF2, 0xF0))
add_rect(s3, LEFT_MARGIN, Inches(2.95), Inches(0.06), Inches(1.55), TEAL)
add_text(s3, "Step one, answer a few quick questions", LEFT_MARGIN + Inches(0.25), Inches(3.12), LEFT_COL_W - Inches(0.5), Inches(0.35),
          font_name=FONT_BODY, font_size=13, bold=True, colour=BLACK)
add_text(s3, "See the solutions matched to your size and industry.", LEFT_MARGIN + Inches(0.25), Inches(3.5), LEFT_COL_W - Inches(0.5), Inches(0.3),
          font_name=FONT_BODY, font_size=11, colour=RGBColor(0x33, 0x33, 0x33))
add_hyperlink_text(s3, "profit-pulse.com.au/full-suite-of-products",
                    "https://profit-pulse.com.au/full-suite-of-products",
                    LEFT_MARGIN + Inches(0.25), Inches(3.85), LEFT_COL_W - Inches(0.5), Inches(0.3),
                    font_size=12, bold=True, colour=TEAL)

# Direct purchase CTA
add_rect(s3, LEFT_MARGIN, Inches(4.7), LEFT_COL_W, Inches(0.65), AMBER_D)
add_hyperlink_text(s3, "Purchase the suggested product now to get started",
                    "https://buy.stripe.com/eVqdRad9e66cftK23Z3ks0h",
                    LEFT_MARGIN + Inches(0.25), Inches(4.9), LEFT_COL_W - Inches(0.5), Inches(0.35),
                    font_size=13, bold=True, colour=BLACK)

# Booking alternative
add_text(s3, "Prefer a conversation first?", LEFT_MARGIN, Inches(5.6), LEFT_COL_W, Inches(0.3),
          font_name=FONT_BODY, font_size=11, colour=RGBColor(0x33, 0x33, 0x33))
add_hyperlink_text(s3, "Book a complimentary discovery call",
                    "https://bookings.cloud.microsoft/book/ProfitPulse1@profit-pulse.com.au/?ismsaljsauthenabled=true",
                    LEFT_MARGIN, Inches(5.92), LEFT_COL_W, Inches(0.3),
                    font_size=12, bold=True, colour=TEAL)

# Right column: About ProfitPulse and Nitesh panel
add_rect(s3, RIGHT_COL_X, Inches(1.15), RIGHT_COL_W, Inches(5.55), BLACK)
add_text(s3, "Nitesh Roopa", RIGHT_COL_X + Inches(0.25), Inches(1.35), RIGHT_COL_W - Inches(0.5), Inches(0.4),
          font_name=FONT_HEAD, font_size=18, bold=True, colour=AMBER_B)
add_text(s3, "CA, Managing Partner, ProfitPulse", RIGHT_COL_X + Inches(0.25), Inches(1.78), RIGHT_COL_W - Inches(0.5), Inches(0.32),
          font_name=FONT_BODY, font_size=12, colour=TEAL)

cred_points = [
    "16 years across 4 countries",
    "52 deals executed and managed",
    "Largest deal USD 1.3 billion, Cahora Bassa",
    "GRBT project value over AUD 10 billion",
]
y = Inches(2.35)
for cp in cred_points:
    add_text(s3, cp, RIGHT_COL_X + Inches(0.25), y, RIGHT_COL_W - Inches(0.5), Inches(0.4),
              font_name=FONT_BODY, font_size=11, colour=OFF_WHITE, line_spacing=1.1)
    y += Inches(0.42)

add_rect(s3, RIGHT_COL_X + Inches(0.25), Inches(4.15), RIGHT_COL_W - Inches(0.5), Pt(1), TEAL)

contact_lines = [
    "Profit-Pulse.com.au",
    "Nitesh@Profit-Pulse.com.au",
    "+61 411 876 267",
    "linkedin.com/in/nitesh-roopa-77594163",
]
y = Inches(4.35)
for cl in contact_lines:
    add_text(s3, cl, RIGHT_COL_X + Inches(0.25), y, RIGHT_COL_W - Inches(0.5), Inches(0.35),
              font_name=FONT_BODY, font_size=11, colour=OFF_WHITE)
    y += Inches(0.38)

out_path = "/home/user/Daily-Outreach/Out-reach efforts/Brief_TruckitNet_02Jul2026.pptx"
prs.save(out_path)

# Patch the theme colour scheme so LibreOffice renders hyperlink runs in a
# brand colour instead of its default blue/purple, per Rule 3.
import zipfile
import shutil

tmp_path = out_path + ".tmp"
with zipfile.ZipFile(out_path, "r") as zin:
    names = zin.namelist()
    theme_xml = zin.read("ppt/theme/theme1.xml").decode("utf-8")
    # Black reads correctly in every context this deck uses a hyperlink in
    # (the amber CTA tile and the white body background), so it is the safe
    # fallback for renderers, such as LibreOffice, that ignore the explicit
    # per run colour set on hyperlinked text.
    theme_xml = theme_xml.replace(
        "<a:hlink><a:srgbClr val=\"0000FF\"/></a:hlink>",
        "<a:hlink><a:srgbClr val=\"000000\"/></a:hlink>",
    ).replace(
        "<a:folHlink><a:srgbClr val=\"800080\"/></a:folHlink>",
        "<a:folHlink><a:srgbClr val=\"000000\"/></a:folHlink>",
    )
    with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in names:
            data = zin.read(item)
            if item == "ppt/theme/theme1.xml":
                data = theme_xml.encode("utf-8")
            zout.writestr(item, data)
shutil.move(tmp_path, out_path)

print(f"PPTX saved: {out_path}")

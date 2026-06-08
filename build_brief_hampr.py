#!/usr/bin/env python3
"""Build Brief_Hampr_09Jun2026.pptx — ProfitPulse Nightly Outreach Engine"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree
from pptx.oxml.ns import qn
import os

# ── Brand colours ─────────────────────────────────────────────────────────────
BLK  = RGBColor(0x00, 0x00, 0x00)
TEAL = RGBColor(0x01, 0xA2, 0x96)
AMB  = RGBColor(0xF8, 0xC8, 0x06)
AMBD = RGBColor(0xF6, 0xA1, 0x02)
GOLD = RGBColor(0xE3, 0xA7, 0x12)
WHT  = RGBColor(0xFF, 0xFF, 0xFF)
OWHT = RGBColor(0xE6, 0xE5, 0xDE)

OUT  = "/home/user/Daily-Outreach/Out-reach efforts"
DATE = "09 Jun 2026"

STRIPE_B1_EXEC    = "https://buy.stripe.com/dRmbJ22uA8ekdlCgYT3ks0r"
QUEST_CLEAN       = "https://profit-pulse.com.au/full-suite-of-products"
BOOKING           = ("https://bookings.cloud.microsoft/book/"
                     "ProfitPulse1@profit-pulse.com.au/?ismsaljsauthenabled=true")

prs = Presentation()
prs.slide_width  = Inches(10)
prs.slide_height = Inches(5.625)
blank = prs.slide_layouts[6]


# ── Helpers ───────────────────────────────────────────────────────────────────

def add_slide():
    return prs.slides.add_slide(blank)


def rect(sl, l, t, w, h, fill, no_line=True):
    s = sl.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    return s


def tb(sl, l, t, w, h, text, sz=10, bold=False, clr=BLK,
       font="Calibri", align=PP_ALIGN.LEFT, italic=False,
       ml=0.04, mt=0.02, url=None):
    box = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = box.text_frame
    tf.word_wrap    = True
    tf.margin_left  = Inches(ml)
    tf.margin_top   = Inches(mt)
    tf.margin_right = Inches(0.02)
    tf.margin_bottom= Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text           = text
    r.font.name      = font
    r.font.size      = Pt(sz)
    r.font.color.rgb = clr
    r.font.bold      = bold
    r.font.italic    = italic
    if url:
        _add_run_hyperlink(r, sl, url)
    return box


def _add_run_hyperlink(run, slide, url):
    r_elem = run._r
    rPr = r_elem.find(qn("a:rPr"))
    if rPr is None:
        rPr = etree.Element(qn("a:rPr"))
        r_elem.insert(0, rPr)
    rel = slide.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hlink = etree.SubElement(rPr, qn("a:hlinkClick"))
    hlink.set(qn("r:id"), rel)


def frame(sl, eyebrow):
    """Amber left stripe, black header band, footer line on every slide."""
    rect(sl, 0, 0, 0.08, 5.625, AMB)
    rect(sl, 0.08, 0, 9.92, 0.85, BLK)
    tb(sl, 0.2, 0.07, 5.8, 0.38, eyebrow,
       sz=9.5, bold=True, clr=OWHT, font="Calibri")
    tb(sl, 5.5, 0.07, 4.4, 0.38, "PROFITPULSE",
       sz=9.5, bold=True, clr=OWHT, font="Calibri", align=PP_ALIGN.RIGHT)
    rect(sl, 0.08, 5.47, 9.92, 0.015, OWHT)
    tb(sl, 0.18, 5.49, 9.55, 0.13,
       f"Prepared by Nitesh Roopa CA, Managing Partner, ProfitPulse  |  Profit-Pulse.com.au  |  {DATE}",
       sz=6.5, italic=True, clr=BLK)


def stat_card(sl, l, t, w, h, num, lab1, lab2, src):
    rect(sl, l, t, w, h, BLK)
    rect(sl, l, t, w, 0.045, TEAL)
    tb(sl, l+0.05, t+0.07, w-0.1, h*0.41,
       num, sz=19, bold=True, clr=AMB, font="Georgia")
    tb(sl, l+0.05, t+h*0.48, w-0.1, h*0.20, lab1, sz=7.5, clr=WHT)
    tb(sl, l+0.05, t+h*0.66, w-0.1, h*0.18, lab2, sz=7.5, clr=OWHT)
    tb(sl, l+0.05, t+h*0.83, w-0.1, h*0.15, src, sz=6, clr=TEAL, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Commercial Intelligence Brief
# ══════════════════════════════════════════════════════════════════════════════
s1 = add_slide()
frame(s1, "COMMERCIAL INTELLIGENCE BRIEF")

tb(s1, 0.2, 0.88, 7, 0.52, "Hampr",
   sz=36, bold=True, clr=BLK, font="Georgia")
tb(s1, 0.2, 1.37, 9.6, 0.22,
   "B2B workplace hospitality and corporate food technology marketplace   |   Sydney, NSW",
   sz=10, clr=BLK)

# 6 stat cards  y=1.62, h=1.07
CW = 1.565; CH = 1.07; CY = 1.62; GAP = 0.064
XS = [0.12 + i * (CW + GAP) for i in range(6)]

CARDS = [
    ("AUD $12.7M",  "Revenue",               "FY2024",               "Smart50 2024, SmartCompany Nov 2024"),
    ("153%",        "Revenue growth",         "year on year FY2024",  "Smart50 2024, SmartCompany Nov 2024"),
    ("#8",          "Smart50 ranking",        "nationally, 2024",     "SmartCompany Smart50 Nov 2024"),
    ("37",          "Employees",              "full time equivalents","Smart50 2024, SmartCompany Nov 2024"),
    ("2018",        "Year founded",           "Sydney, NSW",          "SmartCompany Smart50 2024 profile"),
    ("$11.1M",      "Total funding raised",   "across 4 rounds",      "Crunchbase / Tracxn, May 2025"),
]
for i, (num, l1, l2, src) in enumerate(CARDS):
    stat_card(s1, XS[i], CY, CW, CH, num, l1, l2, src)

# Key commercial signals
tb(s1, 0.2, 2.76, 4.5, 0.19,
   "KEY COMMERCIAL SIGNALS", sz=8.5, bold=True, clr=TEAL)

SIGNALS = [
    "Singapore market entry FY2026: hampr.sg website live, international expansion stated publicly.  Source: Smart50 2024 profile, Nov 2024",
    "Two major enterprise tenders targeted within 12 months of Smart50 submission.  Source: Smart50 2024 profile, Nov 2024",
    "Enterprise client publicly named: Toyota Finance Australia referenced in company media.  Source: Momentum91 podcast, May 2025",
    "Investor backing includes Blackbird Ventures, Techstars, Startmate, Lee Kim Tah Group.  Source: Crunchbase / Tracxn, 2025",
    "National operations active across Sydney, Melbourne, Perth, Brisbane, and Canberra.  Source: Hampr website, 2025",
    "No publicly named CFO or Head of Finance in any company profile or media coverage.  Source: Multiple profiles, 2025",
]
SY = 2.97
for sig in SIGNALS:
    tb(s1, 0.2, SY, 9.6, 0.34, sig, sz=8.5, clr=BLK)
    SY += 0.345


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — The Opportunity
# ══════════════════════════════════════════════════════════════════════════════
s2 = add_slide()
frame(s2, "THE OPPORTUNITY")

tb(s2, 0.2, 0.89, 9.5, 0.3,
   "Hampr   |   Three commercial observations from ProfitPulse",
   sz=12, bold=True, clr=BLK, font="Georgia")

# Three columns: teal / black / gold
COLS = [
    (0.10, 3.22, RGBColor(0x01, 0xA2, 0x96), WHT, WHT, BLK),
    (3.45, 3.22, BLK,                          WHT, WHT, AMB),
    (6.80, 3.10, GOLD,                          BLK, BLK, BLK),
]
IDXS  = ["01", "02", "03"]
HEADS = [
    "The marketplace float is a cash flow event waiting to happen",
    "The Singapore launch needs a financial model, not just a website",
    "Enterprise contracts will change the financial DNA of the business",
]
BODIES = [
    ("At $12.7M in revenue growing at 153% year on year, Hampr processes payments between "
     "corporate clients and a network of food and event suppliers. As transaction volume grows "
     "so does the cash sitting in transit between client receipts and supplier payments. "
     "Without a structured 13 week cash flow forecast and a treasury playbook a fast growing "
     "marketplace can run short of cash even while reporting a healthy profit. "
     "ProfitPulse fractional CFO engagements routinely identify six to twelve percent of "
     "annual revenue sitting in manageable float that can be put to work."),

    ("Hampr has moved into Singapore with a live website and a publicly stated FY2026 market "
     "entry target. International expansion requires upfront investment in team, regulatory "
     "compliance, and market development well ahead of local revenue. With the most recent "
     "funding round at $1.08M in May 2025, a rigorous financial model for the Singapore "
     "launch is needed: one that maps cash burn, the revenue ramp, and the break even "
     "timeline. Building and maintaining that model within a monthly reporting rhythm is "
     "exactly what a Fractional CFO Partnership is designed to carry."),

    ("Two enterprise tenders in active pursuit means a step change from a diversified SME "
     "client base to a smaller number of much larger contracts. Enterprise clients typically "
     "carry 60 to 90 day payment terms, complex invoicing, and detailed financial reporting "
     "expectations. This concentrates revenue risk, extends the cash conversion cycle, and "
     "demands board grade financial reporting that a founder led business at this scale may "
     "not yet have in place. Getting ahead of that transition now protects both margin and "
     "negotiating position when the contracts land."),
]

CT = 1.24; CH2 = 3.97
for (cx, cw, bg, tc, hc, nc), idx, head, body in zip(COLS, IDXS, HEADS, BODIES):
    rect(s2, cx, CT, cw, CH2, bg)
    tb(s2, cx+0.1, CT+0.09, cw-0.2, 0.44, idx, sz=26, bold=True, clr=nc, font="Georgia")
    tb(s2, cx+0.1, CT+0.55, cw-0.2, 0.6,  head, sz=9, bold=True, clr=hc)
    tb(s2, cx+0.1, CT+1.17, cw-0.2, 2.73, body, sz=8.5, clr=tc)

tb(s2, 0.15, 5.27, 9.65, 0.18,
   ("These observations are offered in good faith. Hampr has built something genuinely "
    "impressive. The question is simply whether the financial architecture is ready for the next chapter."),
   sz=8, italic=True, clr=BLK)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — The Recommendation
# ══════════════════════════════════════════════════════════════════════════════
s3 = add_slide()
frame(s3, "THE RECOMMENDATION")

# ── Left column ───────────────────────────────────────────────────────────────
LX = 0.12; LW = 5.55

tb(s3, LX+0.05, 0.9, LW, 0.5,
   "Fractional CFO Partnership",
   sz=21, bold=True, clr=BLK, font="Georgia")
tb(s3, LX+0.05, 1.38, LW, 0.28,
   "$7,500 per month",
   sz=13, bold=True, clr=TEAL)
tb(s3, LX+0.05, 1.66, LW, 0.78,
   ("A senior financial partner at the table on a monthly cadence. Includes monthly management "
    "pack, quarterly board grade review, ad hoc decision support, and a single annual deep dive. "
    "For Hampr this means a structured financial rhythm for the Singapore launch, enterprise "
    "contract cash flow modelling, and the reporting infrastructure to match the ambition."),
   sz=9.5, clr=BLK)

rect(s3, LX+0.05, 2.47, LW-0.05, 0.015, TEAL)

# Step one block
rect(s3, LX, 2.5, 0.04, 1.42, TEAL)
tb(s3, LX+0.12, 2.53, LW-0.05, 0.22, "Step one: answer a few quick questions",
   sz=10, bold=True, clr=BLK)
tb(s3, LX+0.12, 2.75, LW-0.05, 0.2,
   "See the solutions matched to your size and industry.",
   sz=9.5, clr=BLK)
tb(s3, LX+0.12, 2.95, LW-0.05, 0.22,
   "profit-pulse.com.au/full-suite-of-products",
   sz=9.5, bold=True, clr=TEAL, url=QUEST_CLEAN)

rect(s3, LX+0.05, 3.22, LW-0.05, 0.015, AMBD)

# Direct purchase CTA
rect(s3, LX, 3.24, 0.04, 0.54, AMBD)
tb(s3, LX+0.12, 3.27, LW-0.05, 0.22,
   "Already know this is the priority?",
   sz=9.5, bold=True, clr=BLK)
tb(s3, LX+0.12, 3.49, LW-0.05, 0.22,
   "Purchase the suggested product now to get started",
   sz=9.5, bold=True, clr=AMBD, url=STRIPE_B1_EXEC)

rect(s3, LX+0.05, 3.82, LW-0.05, 0.015, TEAL)

# Discovery call
rect(s3, LX, 3.85, 0.04, 0.56, AMB)
tb(s3, LX+0.12, 3.87, LW-0.05, 0.22, "Prefer a conversation first?",
   sz=9.5, bold=True, clr=BLK)
tb(s3, LX+0.12, 4.09, LW-0.05, 0.3,
   "Book a complimentary discovery call with Nitesh Roopa",
   sz=9.5, clr=TEAL, url=BOOKING)

# ── Vertical divider ─────────────────────────────────────────────────────────
rect(s3, 5.75, 0.88, 0.03, 4.56, TEAL)

# ── Right column — credibility panel ─────────────────────────────────────────
RX = 5.83; RW = 4.05
rect(s3, RX, 0.88, RW, 4.56, BLK)

tb(s3, RX+0.12, 0.95, RW-0.2, 0.46, "Nitesh Roopa",
   sz=18, bold=True, clr=AMB, font="Georgia")
tb(s3, RX+0.12, 1.39, RW-0.2, 0.24,
   "CA, Managing Partner   |   ProfitPulse",
   sz=10, clr=OWHT)

rect(s3, RX+0.12, 1.65, RW-0.3, 0.02, TEAL)

CREDS = [
    "16 years across 4 countries",
    "52 deals executed and managed across the career",
    "Largest single deal: USD 1.3 billion, Cahora Bassa, Mozambique Government, Hydro",
    "Total GRBT project value in Queensland: over AUD 10 billion",
]
CY3 = 1.72
for c in CREDS:
    tb(s3, RX+0.12, CY3, RW-0.2, 0.32, c, sz=8.5, clr=OWHT)
    CY3 += 0.32

rect(s3, RX+0.12, CY3+0.04, RW-0.3, 0.02, TEAL)
CY3 += 0.1

CONTACTS = [
    ("Profit-Pulse.com.au",                    QUEST_CLEAN),
    ("Nitesh@Profit-Pulse.com.au",             None),
    ("+61 411 876 267",                        None),
    ("linkedin.com/in/nitesh-roopa-77594163",  "https://linkedin.com/in/nitesh-roopa-77594163"),
]
for (ctxt, curl) in CONTACTS:
    tb(s3, RX+0.12, CY3, RW-0.2, 0.28, ctxt, sz=8.5, clr=TEAL, url=curl)
    CY3 += 0.28

# ── Save ─────────────────────────────────────────────────────────────────────
out_path = os.path.join(OUT, "Brief_Hampr_09Jun2026.pptx")
prs.save(out_path)
print(f"Saved: {out_path}")

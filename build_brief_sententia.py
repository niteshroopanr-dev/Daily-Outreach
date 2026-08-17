"""
ProfitPulse Brief Builder
Target: Sententia Consulting | Date: 18 Aug 2026
Three slide prospect facing deck. House style per Section 6, v3.3.
Brand colours only. Zero dashes.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

# Brand colours, the seven only
BLACK     = RGBColor(0x00, 0x00, 0x00)
TEAL      = RGBColor(0x01, 0xA2, 0x96)
AMBER_B   = RGBColor(0xF8, 0xC8, 0x06)
AMBER_D   = RGBColor(0xF6, 0xA1, 0x02)
GOLD      = RGBColor(0xE3, 0xA7, 0x12)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xE6, 0xE5, 0xDE)

SERIF = "Georgia"
SANS  = "Calibri"

W = Inches(13.333)
H = Inches(7.5)

DATE_STAMP = "18 Aug 2026"
COMPANY = "Sententia Consulting"


def set_bg(slide, colour):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = colour


def add_rect(slide, left, top, width, height, colour, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = colour
    if not line:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = colour
    shp.shadow.inherit = False
    return shp


def add_text(slide, text, left, top, width, height, font=SANS, size=12,
             bold=False, italic=False, colour=BLACK, align=PP_ALIGN.LEFT,
             anchor=None, hyperlink=None, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = colour
    if hyperlink:
        run.hyperlink.address = hyperlink
    return box


def add_multiline(slide, lines, left, top, width, height, font=SANS,
                   default_size=12, default_colour=BLACK, align=PP_ALIGN.LEFT,
                   space_after=4, anchor=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass
    if anchor is not None:
        tf.vertical_anchor = anchor
    first = True
    for line in lines:
        if isinstance(line, str):
            cfg = {"text": line, "size": default_size, "colour": default_colour,
                   "bold": False, "italic": False, "font": font}
        else:
            cfg = {"text": line.get("text", ""), "size": line.get("size", default_size),
                   "colour": line.get("colour", default_colour), "bold": line.get("bold", False),
                   "italic": line.get("italic", False), "font": line.get("font", font)}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = cfg["text"]
        run.font.name = cfg["font"]
        run.font.size = Pt(cfg["size"])
        run.font.bold = cfg["bold"]
        run.font.italic = cfg["italic"]
        run.font.color.rgb = cfg["colour"]
    return box


def add_chrome(slide, eyebrow, header_right="PROFITPULSE"):
    """Left accent stripe, black header band, footer line. Fixed chrome."""
    add_rect(slide, Inches(0), Inches(0), Inches(0.1), H, AMBER_D)
    add_rect(slide, Inches(0), Inches(0), W, Inches(1.0), BLACK)
    add_text(slide, eyebrow, Inches(0.35), Inches(0.32), Inches(8.0), Inches(0.4),
              font=SANS, size=12, bold=True, colour=OFF_WHITE, align=PP_ALIGN.LEFT)
    add_text(slide, header_right, Inches(8.5), Inches(0.32), Inches(4.5), Inches(0.4),
              font=SANS, size=13, bold=True, colour=AMBER_B, align=PP_ALIGN.RIGHT)
    # Footer
    add_text(slide, "Prepared by Nitesh Roopa CA, Managing Partner, ProfitPulse, Profit-Pulse.com.au",
              Inches(0.35), Inches(7.08), Inches(9.5), Inches(0.3),
              font=SANS, size=8, bold=False, colour=BLACK, align=PP_ALIGN.LEFT)
    add_text(slide, DATE_STAMP, Inches(10.0), Inches(7.08), Inches(2.9), Inches(0.3),
              font=SANS, size=8, bold=False, colour=BLACK, align=PP_ALIGN.RIGHT)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]

# ════════════════════════════════════════════════════════════════════════
# SLIDE 1, THE COMMERCIAL INTELLIGENCE BRIEF
# ════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank)
set_bg(s1, WHITE)
add_chrome(s1, "COMMERCIAL INTELLIGENCE BRIEF")

add_text(s1, COMPANY, Inches(0.35), Inches(1.15), Inches(10.5), Inches(0.85),
          font=SERIF, size=42, bold=True, colour=BLACK, align=PP_ALIGN.LEFT)

add_text(s1, "Government risk and financial advisory firm, Barton, Canberra ACT",
          Inches(0.35), Inches(2.0), Inches(11.5), Inches(0.4),
          font=SANS, size=14, bold=False, colour=TEAL, align=PP_ALIGN.LEFT)

# Stat cards, five, even row, top y 2.4, height 1.6
CARD_Y = Inches(2.4)
CARD_H = Inches(1.6)
CARD_W = Inches(2.3)
CARD_GAP = Inches(0.15)
CARD_START_X = Inches(0.35)

stat_cards = [
    ("$10M+", ["Revenue surpassed,", "FY2025"], "AFR Fast 100 2025"),
    ("60.2%", ["Three year revenue", "CAGR"], "AFR Fast 100, rank 45"),
    ("2020", ["Founded as a two", "person firm"], "Consultancy.com.au"),
    ("$45.3M", ["Lifetime federal", "contract value"], "Pollywatch, FY20 to FY26"),
    ("174", ["Federal contract", "notices won"], "Pollywatch register"),
]

x = CARD_START_X
for number, label_lines, source in stat_cards:
    add_rect(s1, x, CARD_Y, CARD_W, CARD_H, BLACK)
    add_rect(s1, x, CARD_Y, CARD_W, Inches(0.06), TEAL)
    add_text(s1, number, x + Inches(0.15), CARD_Y + Inches(0.14), CARD_W - Inches(0.3), Inches(0.5),
              font=SERIF, size=28, bold=True, colour=AMBER_B, align=PP_ALIGN.LEFT)
    add_multiline(s1, label_lines, x + Inches(0.15), CARD_Y + Inches(0.68), CARD_W - Inches(0.3), Inches(0.55),
                   font=SANS, default_size=11, default_colour=OFF_WHITE, space_after=0)
    add_text(s1, source, x + Inches(0.15), CARD_Y + Inches(1.32), CARD_W - Inches(0.3), Inches(0.24),
              font=SANS, size=8, italic=True, colour=OFF_WHITE, align=PP_ALIGN.LEFT)
    x += CARD_W + CARD_GAP

# Key commercial signals
add_text(s1, "KEY COMMERCIAL SIGNALS", Inches(0.35), Inches(4.28), Inches(6), Inches(0.3),
          font=SANS, size=12, bold=True, colour=TEAL, align=PP_ALIGN.LEFT)

signals = [
    "Melbourne office opened Aug 2024, led by new state MD Tom Fazio (Consultancy.com.au)",
    "Brioni Bale promoted to Managing Director, four senior hires added in 2025 (Consultancy.com.au)",
    "Two agencies are over half of the $45.3M lifetime federal contract book (Pollywatch)",
    "FY2025 26 alone brought $7.15M in new Commonwealth contract awards (Pollywatch)",
    "Grew from 2 founders in 2020 to a near 20 strong team by 2023 (Consultancy.com.au)",
]
sig_lines = [{"text": "•  " + s, "size": 12, "colour": BLACK} for s in signals]
add_multiline(s1, sig_lines, Inches(0.35), Inches(4.62), Inches(12.4), Inches(2.3),
               font=SANS, default_size=12, default_colour=BLACK, space_after=10)

prs_slide1 = s1

# ════════════════════════════════════════════════════════════════════════
# SLIDE 2, THE OPPORTUNITY
# ════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank)
set_bg(s2, WHITE)
add_chrome(s2, "THE OPPORTUNITY")

add_text(s2, "Sententia Consulting: three commercial observations from ProfitPulse",
          Inches(0.35), Inches(1.18), Inches(12.4), Inches(0.45),
          font=SERIF, size=18, bold=True, colour=BLACK, align=PP_ALIGN.LEFT)

COL_Y = Inches(1.75)
COL_H = Inches(4.5)
COL_W = Inches(4.1667)
COL_GAP = Inches(0.1)
COL_START_X = Inches(0.3)

columns = [
    {
        "fill": TEAL, "text_colour": BLACK, "index": "01",
        "header": "Growth is outrunning infrastructure",
        "body": ("A 60.2 percent three year revenue CAGR took Sententia from two "
                 "founders to a near 20 strong team and a second state office in "
                 "under five years, per the 2025 AFR Fast 100 and Consultancy.com.au. "
                 "At this pace, capacity planning and capital allocation typically lag "
                 "delivery. A Strategic Growth Diagnostic maps revenue, capacity and "
                 "margin headroom into a costed 12 month plan before that gap widens."),
    },
    {
        "fill": BLACK, "text_colour": WHITE, "index": "02",
        "header": "Two agencies carry over half the book",
        "body": ("Public federal contract records show Services Australia and the "
                 "Department of Health and Aged Care together account for more than "
                 "half of Sententia's $45.3 million lifetime Commonwealth contract "
                 "value, per Pollywatch. A strong panel position, and a concentration "
                 "worth watching as a standing item, the kind of visibility a "
                 "Fractional CFO Partnership keeps in view every month."),
    },
    {
        "fill": GOLD, "text_colour": BLACK, "index": "03",
        "header": "Two offices means two engines to plan",
        "body": ("Canberra and Melbourne each carry their own delivery leadership "
                 "and their own billable mix. As headcount keeps growing across both "
                 "sites, per Consultancy.com.au reporting on recent hires, seeing "
                 "utilisation and revenue per consultant clearly across two offices "
                 "gets harder without dedicated tracking. A Workforce Capacity and "
                 "Utilisation Review builds that view."),
    },
]

cx = COL_START_X
for col in columns:
    add_rect(s2, cx, COL_Y, COL_W, COL_H, col["fill"])
    add_text(s2, col["index"], cx + Inches(0.25), COL_Y + Inches(0.15), COL_W - Inches(0.5), Inches(0.85),
              font=SERIF, size=44, bold=True, colour=col["text_colour"], align=PP_ALIGN.LEFT)
    add_text(s2, col["header"], cx + Inches(0.25), COL_Y + Inches(1.05), COL_W - Inches(0.5), Inches(0.65),
              font=SERIF, size=15, bold=True, colour=col["text_colour"], align=PP_ALIGN.LEFT)
    add_text(s2, col["body"], cx + Inches(0.25), COL_Y + Inches(1.75), COL_W - Inches(0.5), Inches(2.6),
              font=SANS, size=10.5, bold=False, colour=col["text_colour"], align=PP_ALIGN.LEFT)
    cx += COL_W + COL_GAP

add_text(s2, ("These are observations offered in good faith. Sententia has built something "
              "impressive. The question is simply whether the financial architecture matches the ambition."),
          Inches(0.3), Inches(6.38), Inches(12.7), Inches(0.5),
          font=SANS, size=11, italic=True, colour=BLACK, align=PP_ALIGN.CENTER)

# ════════════════════════════════════════════════════════════════════════
# SLIDE 3, THE RECOMMENDATION AND HOW TO START
# ════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank)
set_bg(s3, WHITE)
add_chrome(s3, "THE RECOMMENDATION")

LEFT_X = Inches(0.35)
LEFT_W = Inches(7.4)

add_text(s3, "Strategic Growth Diagnostic", LEFT_X, Inches(1.15), LEFT_W, Inches(0.55),
          font=SERIF, size=26, bold=True, colour=BLACK, align=PP_ALIGN.LEFT)
add_text(s3, "$5,000 one off", LEFT_X, Inches(1.68), LEFT_W, Inches(0.4),
          font=SANS, size=17, bold=True, colour=AMBER_D, align=PP_ALIGN.LEFT)
add_text(s3, ("Six weeks mapping revenue, capacity and margin headroom into a 12 month "
              "growth plan with funding and capital allocation steps, sized to a two office model."),
          LEFT_X, Inches(2.12), LEFT_W, Inches(0.7),
          font=SANS, size=12, bold=False, colour=BLACK, align=PP_ALIGN.LEFT)

# Step one block, teal bordered
add_rect(s3, LEFT_X, Inches(2.95), LEFT_W, Inches(1.55), WHITE, line=True)
step_box = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_X, Inches(2.95), LEFT_W, Inches(1.55))
step_box.fill.solid(); step_box.fill.fore_color.rgb = WHITE
step_box.line.color.rgb = TEAL
step_box.line.width = Pt(1.5)
step_box.shadow.inherit = False

add_text(s3, "STEP ONE, ANSWER A FEW QUICK QUESTIONS", LEFT_X + Inches(0.2), Inches(3.08), LEFT_W - Inches(0.4), Inches(0.3),
          font=SANS, size=11, bold=True, colour=TEAL, align=PP_ALIGN.LEFT)
add_text(s3, "See the solutions matched to your size and industry.", LEFT_X + Inches(0.2), Inches(3.4), LEFT_W - Inches(0.4), Inches(0.3),
          font=SANS, size=11, bold=False, colour=BLACK, align=PP_ALIGN.LEFT)
add_text(s3, "profit-pulse.com.au/services/find-your-fit", LEFT_X + Inches(0.2), Inches(3.75), LEFT_W - Inches(0.4), Inches(0.4),
          font=SANS, size=14, bold=True, colour=TEAL, align=PP_ALIGN.LEFT,
          hyperlink="https://profit-pulse.com.au/services/find-your-fit/")

# Direct purchase CTA, amber outlined button
btn = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LEFT_X, Inches(4.7), Inches(4.6), Inches(0.55))
btn.fill.solid(); btn.fill.fore_color.rgb = WHITE
btn.line.color.rgb = AMBER_D
btn.line.width = Pt(1.5)
btn.shadow.inherit = False
btf = btn.text_frame
btf.word_wrap = True
bp = btf.paragraphs[0]
bp.alignment = PP_ALIGN.CENTER
brun = bp.add_run()
brun.text = "Purchase the suggested product now to get started"
brun.font.name = SANS
brun.font.size = Pt(11)
brun.font.bold = True
brun.font.color.rgb = AMBER_D
brun.hyperlink.address = "https://buy.stripe.com/eVqdRad9e66cftK23Z3ks0h"

add_text(s3, "Prefer a conversation first?", LEFT_X, Inches(5.5), LEFT_W, Inches(0.3),
          font=SANS, size=11, bold=False, colour=BLACK, align=PP_ALIGN.LEFT)
add_text(s3, "Book a complimentary discovery call", LEFT_X, Inches(5.82), LEFT_W, Inches(0.35),
          font=SANS, size=12, bold=True, colour=TEAL, align=PP_ALIGN.LEFT,
          hyperlink="https://bookings.cloud.microsoft/book/ProfitPulse1@profit-pulse.com.au/?ismsaljsauthenabled=true")

# Right column, credibility panel
RIGHT_X = Inches(8.05)
RIGHT_W = Inches(4.95)
add_rect(s3, RIGHT_X, Inches(1.15), RIGHT_W, Inches(5.55), BLACK)
add_rect(s3, RIGHT_X, Inches(1.15), RIGHT_W, Inches(0.06), TEAL)

add_text(s3, "Nitesh Roopa", RIGHT_X + Inches(0.25), Inches(1.35), RIGHT_W - Inches(0.5), Inches(0.45),
          font=SERIF, size=19, bold=True, colour=AMBER_B, align=PP_ALIGN.LEFT)
add_text(s3, "CA, Managing Partner, ProfitPulse", RIGHT_X + Inches(0.25), Inches(1.82), RIGHT_W - Inches(0.5), Inches(0.35),
          font=SANS, size=12, bold=False, colour=WHITE, align=PP_ALIGN.LEFT)

cred_lines = [
    "16 years across 4 countries",
    "Over 52 deals executed and managed",
    "Largest single deal USD 1.3 billion, Cahora Bassa",
    "Total GRBT project value over AUD 10 billion",
]
cred = [{"text": "•  " + c, "size": 11, "colour": OFF_WHITE} for c in cred_lines]
add_multiline(s3, cred, RIGHT_X + Inches(0.25), Inches(2.35), RIGHT_W - Inches(0.5), Inches(1.7),
               font=SANS, default_size=11, default_colour=OFF_WHITE, space_after=8)

add_rect(s3, RIGHT_X + Inches(0.25), Inches(4.15), RIGHT_W - Inches(0.5), Pt(1.2), TEAL)

contact_lines = [
    {"text": "Profit-Pulse.com.au", "size": 11, "colour": OFF_WHITE},
    {"text": "Nitesh@Profit-Pulse.com.au", "size": 11, "colour": TEAL},
    {"text": "+61 411 876 267", "size": 11, "colour": OFF_WHITE},
    {"text": "linkedin.com/in/nitesh-roopa-77594163", "size": 11, "colour": TEAL},
]
add_multiline(s3, contact_lines, RIGHT_X + Inches(0.25), Inches(4.35), RIGHT_W - Inches(0.5), Inches(1.5),
               font=SANS, default_size=11, default_colour=OFF_WHITE, space_after=8)

out_path = "/home/user/Daily-Outreach/Out-reach efforts/Brief_SententiaConsulting_18Aug2026.pptx"
prs.save(out_path)
print(f"PPTX saved: {out_path}")

"""
ProfitPulse Brief Builder
Target: National Media | Date: 13 Jun 2026
Three-slide prospect-facing deck. Brand colours only. Zero dashes.
Section 6 house style: white body, amber left stripe, black header band.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

# Brand colours
BLACK     = RGBColor(0x00, 0x00, 0x00)
TEAL      = RGBColor(0x01, 0xA2, 0x96)
AMBER_B   = RGBColor(0xF8, 0xC8, 0x06)
AMBER_D   = RGBColor(0xF6, 0xA1, 0x02)
GOLD      = RGBColor(0xE3, 0xA7, 0x12)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xE6, 0xE5, 0xDE)
MID_GREY  = RGBColor(0x88, 0x88, 0x88)
DRK_GREY  = RGBColor(0x44, 0x44, 0x44)

W = Inches(13.333)
H = Inches(7.5)

QUESTIONNAIRE_URL = "https://profit-pulse.com.au/full-suite-of-products"
STRIPE_URL        = "https://buy.stripe.com/28EbJ2edi0LSbdugYT3ks1D"
BOOKING_URL       = "https://bookings.cloud.microsoft/book/ProfitPulse1@profit-pulse.com.au/?ismsaljsauthenabled=true"


def set_background(slide, colour):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = colour


def add_rect(slide, left, top, width, height, fill_colour, line_colour=None, line_width=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_colour
    if line_colour:
        shape.line.color.rgb = line_colour
        if line_width:
            shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, left, top, width, height,
             font_name="Arial", font_size=12, bold=False,
             colour=WHITE, align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = colour
    return txBox


def add_multiline(slide, lines, left, top, width, height,
                  font_name="Arial", default_size=12,
                  default_colour=OFF_WHITE, default_bold=False,
                  align=PP_ALIGN.LEFT, spacing_after=None):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        if isinstance(line, str):
            cfg = {"text": line, "size": default_size,
                   "colour": default_colour, "bold": default_bold, "italic": False}
        else:
            cfg = {"text": line.get("text", ""),
                   "size": line.get("size", default_size),
                   "colour": line.get("colour", default_colour),
                   "bold": line.get("bold", default_bold),
                   "italic": line.get("italic", False)}
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = align
        if spacing_after:
            p.space_after = Pt(spacing_after)
        run = p.add_run()
        run.text = cfg["text"]
        run.font.name = font_name
        run.font.size = Pt(cfg["size"])
        run.font.bold = cfg["bold"]
        run.font.italic = cfg["italic"]
        run.font.color.rgb = cfg["colour"]
    return txBox


def hyperlink_run(run, slide, url):
    rId = slide.part.relate_to(url, RT.HYPERLINK, is_external=True)
    rPr = run._r.get_or_add_rPr()
    hlink = etree.SubElement(rPr, qn("a:hlinkClick"))
    hlink.set(qn("r:id"), rId)


def add_linked_text(slide, text, url, left, top, width, height,
                    font_name="Arial", font_size=12, bold=False,
                    colour=TEAL, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = colour
    hyperlink_run(run, slide, url)
    return txBox


def add_stat_card(slide, left, top, w, h, number, label1, label2, source):
    """Black tile with teal top stripe, large number, two-line label, source."""
    add_rect(slide, left, top, w, h, BLACK)
    add_rect(slide, left, top, w, Inches(0.07), TEAL)
    add_text(slide, number,
             left + Inches(0.12), top + Inches(0.1), w - Inches(0.2), Inches(0.42),
             font_size=22, bold=True, colour=OFF_WHITE)
    if label1:
        add_text(slide, label1,
                 left + Inches(0.12), top + Inches(0.54), w - Inches(0.2), Inches(0.22),
                 font_size=10, colour=OFF_WHITE)
    if label2:
        add_text(slide, label2,
                 left + Inches(0.12), top + Inches(0.74), w - Inches(0.2), Inches(0.22),
                 font_size=10, colour=MID_GREY)
    add_text(slide, source,
             left + Inches(0.12), top + Inches(0.95), w - Inches(0.2), Inches(0.18),
             font_size=7, colour=TEAL, italic=True)


def add_header_band(slide, eyebrow, company):
    """Shared black header band with eyebrow label and company name."""
    add_rect(slide, Inches(0.1), Inches(0), W - Inches(0.1), Inches(0.95), BLACK)
    add_text(slide, eyebrow,
             Inches(0.3), Inches(0.18), Inches(8), Inches(0.45),
             font_size=11, bold=True, colour=OFF_WHITE)
    add_text(slide, company,
             Inches(9.0), Inches(0.18), Inches(4.1), Inches(0.45),
             font_size=13, bold=True, colour=TEAL, align=PP_ALIGN.RIGHT)


def add_footer(slide, date_str="13 Jun 2026"):
    add_text(slide,
             "Prepared by Nitesh Roopa CA, Managing Partner and Founder, "
             "ProfitPulse, Profit-Pulse.com.au",
             Inches(0.3), Inches(7.1), Inches(9.5), Inches(0.35),
             font_size=9, colour=DRK_GREY)
    add_text(slide, date_str,
             Inches(10.2), Inches(7.1), Inches(2.9), Inches(0.35),
             font_size=9, colour=DRK_GREY, align=PP_ALIGN.RIGHT)


# ============================================================
# BUILD PRESENTATION
# ============================================================
prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]


# ------------------------------------------------------------
# SLIDE 1: COMMERCIAL INTELLIGENCE BRIEF
# ------------------------------------------------------------
s1 = prs.slides.add_slide(blank)
set_background(s1, WHITE)

# Amber left stripe
add_rect(s1, Inches(0), Inches(0), Inches(0.1), H, AMBER_D)

# Header band
add_header_band(s1, "COMMERCIAL INTELLIGENCE BRIEF", "NATIONAL MEDIA")

# Company name large
add_text(s1, "NATIONAL MEDIA",
         Inches(0.3), Inches(1.02), Inches(9.5), Inches(0.85),
         font_name="Georgia", font_size=40, bold=True, colour=BLACK)

# Descriptor
add_text(s1, "B2B Trade Exhibitions and Events Producer, Bundall, Gold Coast QLD",
         Inches(0.3), Inches(1.87), Inches(11), Inches(0.38),
         font_size=13, colour=TEAL)

# Stat cards (6 cards across)
CARD_W = Inches(2.05)
CARD_H = Inches(1.15)
CARD_Y = Inches(2.33)
GAP    = Inches(0.09)
card_data = [
    ("$18.4M",   "Revenue",        "FY2025",                  "Smart50 2025 award citation"),
    ("42%",      "Three Year",     "Average Growth",           "Smart50 2025 award citation"),
    ("Rank 28",  "Smart50",        "Australia 2025",           "SmartCompany Smart50 2025"),
    ("48",       "Team Members",   "",                         "Smart50 2025 award citation"),
    ("9+",       "Live Event",     "Brands",                   "nationalmedia.com.au"),
    ("550+",     "Combined WHS",   "Show Exhibitors",          "nationalmedia.com.au"),
]
cx = Inches(0.3)
for num, l1, l2, src in card_data:
    add_stat_card(s1, cx, CARD_Y, CARD_W, CARD_H, num, l1, l2, src)
    cx += CARD_W + GAP

# Key Commercial Signals section
SIG_Y = Inches(3.62)
add_text(s1, "KEY COMMERCIAL SIGNALS",
         Inches(0.3), SIG_Y, Inches(12), Inches(0.3),
         font_size=11, bold=True, colour=TEAL)
add_rect(s1, Inches(0.3), SIG_Y + Inches(0.3), Inches(12.6), Pt(1.2), TEAL)

signals = [
    ("FutureBuild Australia launched 11 to 13 June 2026 at ICC Sydney: new event brand debut confirmed",
     "nationalmedia.com.au"),
    ("Mark Harvey presented at SISO CEO Summit on Organic Growth via Launches: expansion strategy from the founder confirmed",
     "LinkedIn public profile, public search results"),
    ("Workplace Health and Safety Show runs three state editions: 550 plus exhibitors across NSW, QLD, and VIC editions combined",
     "nationalmedia.com.au"),
    ("Food and Hospitality Week bundles four separate trade shows under one venue: Restaurant and Foodservice Show, Pizza Pasta and Italian Food Show, Cafe and Coffee Show, and RESTECH",
     "nationalmedia.com.au"),
    ("Smart50 2025 rank 28 of 50: revenue independently verified at $18.4 million, 42 percent three year average growth",
     "SmartCompany Smart50 2025 award citation"),
    ("Founded 1993, over three decades of operations from Bundall, Gold Coast QLD, serving multiple Australian industries",
     "nationalmedia.com.au"),
]

sy = SIG_Y + Inches(0.38)
for sig_text, sig_src in signals:
    add_multiline(s1, [
        {"text": sig_text, "size": 10, "colour": BLACK, "bold": False},
        {"text": "Source: " + sig_src, "size": 7, "colour": MID_GREY, "italic": True},
    ], Inches(0.3), sy, Inches(12.6), Inches(0.4),
       default_size=10, default_colour=BLACK, spacing_after=0)
    sy += Inches(0.49)

add_footer(s1)


# ------------------------------------------------------------
# SLIDE 2: THE OPPORTUNITY
# ------------------------------------------------------------
s2 = prs.slides.add_slide(blank)
set_background(s2, WHITE)

add_rect(s2, Inches(0), Inches(0), Inches(0.1), H, AMBER_D)
add_header_band(s2, "THE OPPORTUNITY", "NATIONAL MEDIA")

sub = ("National Media  —  Three commercial observations from ProfitPulse")
# Note: em dash is allowed in Python string but not in output per Rule 1.
# Use a colon+space instead.
add_text(s2, "National Media: Three commercial observations from ProfitPulse",
         Inches(0.3), Inches(1.0), Inches(12.8), Inches(0.35),
         font_size=12, colour=DRK_GREY, italic=True)

# Three observation columns
COL_W = (W - Inches(0.1)) / 3
COL_Y = Inches(1.38)
COL_H = Inches(5.2)
col_fills = [TEAL, BLACK, GOLD]
col_x     = [Inches(0.1), Inches(0.1) + COL_W, Inches(0.1) + 2 * COL_W]
txt_col   = [BLACK, OFF_WHITE, BLACK]   # Rule 3: black on teal/amber, white on black
src_col   = [RGBColor(0x11, 0x55, 0x50),
             MID_GREY,
             RGBColor(0x55, 0x40, 0x00)]

obs = [
    {
        "index": "01",
        "header": "Nine event brands, one financial picture",
        "para": (
            "National Media runs nine or more distinct event brands across food and "
            "hospitality, workplace health and safety, construction, and accommodation "
            "technology. With $18.4 million in revenue across this portfolio, the "
            "aggregate 42 percent growth headline looks strong. The question that "
            "matters at this stage is whether that growth is distributed evenly or "
            "concentrated in a few events carrying the rest. In most multi brand "
            "businesses, two or three products generate the bulk of the real margin "
            "while others absorb shared overhead without full visibility."
        ),
    },
    {
        "index": "02",
        "header": "Each new launch commits capital before revenue arrives",
        "para": (
            "Mark Harvey's strategy of organic growth via new event launches, stated "
            "publicly at the SISO CEO Summit, means the portfolio is actively expanding. "
            "FutureBuild Australia at ICC Sydney in June 2026 is the most recent addition. "
            "Each launch commits venue deposits, exhibitor acquisition spend, and team time "
            "before the first stand is sold. Without a clear view of which existing events "
            "produce the strongest returns, capital allocation for new launches defaults to "
            "experience and instinct rather than evidence. That is a solvable problem."
        ),
    },
    {
        "index": "03",
        "header": "The next 42 percent requires deliberate capital allocation",
        "para": (
            "With nine or more event brands running across three states and an active "
            "launch strategy, the financial complexity of the portfolio is rising. "
            "Deciding which shows to scale, which to price more assertively, and which "
            "to rationalise requires an objective ranking of every event by gross margin "
            "and contribution. The Product and Service Line Profitability review produces "
            "exactly that: every brand ranked by financial performance in three weeks, "
            "built from management accounts, delivered as a clear action list."
        ),
    },
]

for i, ob in enumerate(obs):
    cx2 = col_x[i]
    add_rect(s2, cx2, COL_Y, COL_W, COL_H, col_fills[i])
    tc = txt_col[i]
    add_text(s2, ob["index"],
             cx2 + Inches(0.18), COL_Y + Inches(0.18), COL_W - Inches(0.3), Inches(0.55),
             font_name="Georgia", font_size=30, bold=True, colour=tc)
    add_text(s2, ob["header"],
             cx2 + Inches(0.18), COL_Y + Inches(0.78), COL_W - Inches(0.3), Inches(0.55),
             font_size=13, bold=True, colour=tc, wrap=True)
    add_text(s2, ob["para"],
             cx2 + Inches(0.18), COL_Y + Inches(1.38), COL_W - Inches(0.3), Inches(3.7),
             font_size=10, colour=tc, wrap=True)

# Warm closing line below columns
add_text(s2,
         "These observations are offered in the spirit of a genuine commercial conversation. "
         "National Media has built something impressive over three decades. "
         "The question is simply whether the financial intelligence behind the portfolio "
         "matches the ambition of the growth strategy.",
         Inches(0.3), Inches(6.65), Inches(12.8), Inches(0.4),
         font_size=10, colour=DRK_GREY, italic=True, align=PP_ALIGN.CENTER)

add_footer(s2)


# ------------------------------------------------------------
# SLIDE 3: THE RECOMMENDATION
# ------------------------------------------------------------
s3 = prs.slides.add_slide(blank)
set_background(s3, WHITE)

add_rect(s3, Inches(0), Inches(0), Inches(0.1), H, AMBER_D)
add_header_band(s3, "THE RECOMMENDATION", "NATIONAL MEDIA")

# LEFT COLUMN: Recommendation and next step
LX  = Inches(0.25)
LW  = Inches(7.9)
LY0 = Inches(1.08)

add_text(s3, "Product and Service Line Profitability",
         LX, LY0, LW, Inches(0.7),
         font_name="Georgia", font_size=22, bold=True, colour=BLACK)

add_text(s3, "$3,950 one off",
         LX, LY0 + Inches(0.72), LW, Inches(0.38),
         font_size=16, bold=True, colour=AMBER_D)

add_text(s3,
         "A three week project ranking every event brand by gross margin, contribution "
         "margin, and operational drag. For a nine brand events business growing at "
         "42 percent, this identifies which shows to scale, which to reprice, and "
         "which to rationalise. Verified ProfitPulse price. "
         "The questionnaire confirms the exact fit for your size and industry.",
         LX, LY0 + Inches(1.14), LW, Inches(0.92),
         font_size=11, colour=BLACK, wrap=True)

# Step one block
add_rect(s3, LX, LY0 + Inches(2.14), LW, Inches(1.6),
         RGBColor(0xF4, 0xF4, 0xF4))
add_text(s3, "STEP ONE: ANSWER A FEW QUICK QUESTIONS",
         LX + Inches(0.15), LY0 + Inches(2.22), LW - Inches(0.3), Inches(0.32),
         font_size=11, bold=True, colour=TEAL)
add_text(s3, "See the solutions matched to your size and industry.",
         LX + Inches(0.15), LY0 + Inches(2.56), LW - Inches(0.3), Inches(0.28),
         font_size=11, colour=BLACK)
# Clean questionnaire address (no UTM on brief)
add_linked_text(s3, "profit-pulse.com.au/full-suite-of-products",
                QUESTIONNAIRE_URL,
                LX + Inches(0.15), LY0 + Inches(2.86), LW - Inches(0.3), Inches(0.3),
                font_size=12, bold=False, colour=TEAL)

# CTA: direct purchase (Stripe link hidden behind these words)
add_rect(s3, LX, LY0 + Inches(3.82), LW, Inches(0.52),
         AMBER_D)
add_linked_text(s3, "Purchase the suggested product now to get started",
                STRIPE_URL,
                LX + Inches(0.2), LY0 + Inches(3.88), LW - Inches(0.4), Inches(0.42),
                font_name="Arial", font_size=13, bold=True, colour=BLACK,
                align=PP_ALIGN.CENTER)

# Conversation alternative
add_text(s3, "Prefer a conversation first?",
         LX, LY0 + Inches(4.44), LW, Inches(0.28),
         font_size=11, colour=BLACK)
add_linked_text(s3, "Book a complimentary discovery call",
                BOOKING_URL,
                LX, LY0 + Inches(4.72), LW, Inches(0.3),
                font_size=11, bold=True, colour=TEAL)

# RIGHT COLUMN: credibility panel
RX = Inches(8.4)
RW = Inches(4.68)
RY = Inches(1.08)
RH = Inches(6.0)

add_rect(s3, RX, RY, RW, RH, BLACK,
         line_colour=RGBColor(0x22, 0x22, 0x22), line_width=Pt(1))

add_text(s3, "NITESH ROOPA",
         RX + Inches(0.2), RY + Inches(0.2), RW - Inches(0.4), Inches(0.5),
         font_size=18, bold=True, colour=AMBER_B)
add_text(s3, "CA, Managing Partner",
         RX + Inches(0.2), RY + Inches(0.72), RW - Inches(0.4), Inches(0.32),
         font_size=12, colour=WHITE)
add_text(s3, "ProfitPulse",
         RX + Inches(0.2), RY + Inches(1.05), RW - Inches(0.4), Inches(0.42),
         font_size=18, bold=True, colour=TEAL)
add_rect(s3, RX + Inches(0.2), RY + Inches(1.5), RW - Inches(0.4), Pt(1.5), TEAL)

contact = [
    {"text": "Profit-Pulse.com.au",          "size": 11, "colour": OFF_WHITE, "bold": False},
    {"text": "Nitesh@Profit-Pulse.com.au",   "size": 11, "colour": TEAL,      "bold": False},
    {"text": "+61 411 876 267",              "size": 11, "colour": OFF_WHITE, "bold": False},
    {"text": "",                              "size": 6,  "colour": WHITE,     "bold": False},
    {"text": "16 years across 4 countries",   "size": 10, "colour": MID_GREY,  "bold": False},
    {"text": "52 deals executed and managed", "size": 10, "colour": MID_GREY,  "bold": False},
    {"text": "Largest single deal USD 1.3 billion Cahora Bassa", "size": 10, "colour": MID_GREY, "bold": False},
    {"text": "Total GRBT project value over AUD 10 billion",     "size": 10, "colour": MID_GREY, "bold": False},
    {"text": "",                              "size": 6,  "colour": WHITE,     "bold": False},
    {"text": "linkedin.com/in/nitesh-roopa-77594163", "size": 9, "colour": DRK_GREY, "bold": False},
]
add_multiline(s3, contact,
              RX + Inches(0.2), RY + Inches(1.58), RW - Inches(0.4), Inches(4.3),
              default_size=11, default_colour=OFF_WHITE, spacing_after=1)

add_footer(s3)


# Save
out_path = "/home/user/Daily-Outreach/Out-reach efforts/Brief_NationalMedia_13Jun2026.pptx"
prs.save(out_path)
print(f"PPTX saved: {out_path}")

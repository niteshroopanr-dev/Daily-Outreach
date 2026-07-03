"""
ProfitPulse Brief Builder, House Style V3.3
Reusable helpers implementing Section 6.0 / 6.0A of the Nightly Outreach Engine spec.
Strict seven colour palette. White body, black header band, amber left stripe,
black stat cards with a teal top edge. No dashes anywhere in generated text.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Seven brand colours, and only these seven, per Rule 3.
BLACK     = RGBColor(0x00, 0x00, 0x00)
TEAL      = RGBColor(0x01, 0xA2, 0x96)
AMBER_B   = RGBColor(0xF8, 0xC8, 0x06)
AMBER_D   = RGBColor(0xF6, 0xA1, 0x02)
GOLD      = RGBColor(0xE3, 0xA7, 0x12)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xE6, 0xE5, 0xDE)

W = Inches(13.333)
H = Inches(7.5)

HEADING_FONT = "Cambria"
BODY_FONT = "Calibri"

CONTENT_TOP = Inches(1.15)
CONTENT_BOTTOM = Inches(6.9)
FOOTER_Y = Inches(7.05)


def new_presentation():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def blank_slide(prs, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def add_rect(slide, left, top, width, height, fill_colour, line_colour=None, line_width=None, shape_type=MSO_SHAPE.RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_colour is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_colour
    if line_colour:
        shape.line.color.rgb = line_colour
        shape.line.width = line_width or Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _apply_run(run, font_name, size, bold, italic, colour, hyperlink=None):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = colour
    if hyperlink:
        run.hyperlink.address = hyperlink


def add_text(slide, text, left, top, width, height, font_name=BODY_FONT, size=12,
             bold=False, italic=False, colour=BLACK, align=PP_ALIGN.LEFT,
             anchor=None, hyperlink=None, line_spacing=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    _apply_run(run, font_name, size, bold, italic, colour, hyperlink)
    return box


def add_multiline(slide, paragraphs, left, top, width, height, align=PP_ALIGN.LEFT,
                   anchor=None, default_font=BODY_FONT):
    """paragraphs: list of dicts {text, size, colour, bold, italic, font, space_after, hyperlink}"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for cfg in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = cfg.get("align", align)
        if cfg.get("space_after") is not None:
            p.space_after = Pt(cfg["space_after"])
        if cfg.get("line_spacing"):
            p.line_spacing = cfg["line_spacing"]
        run = p.add_run()
        run.text = cfg["text"]
        _apply_run(
            run,
            cfg.get("font", default_font),
            cfg.get("size", 12),
            cfg.get("bold", False),
            cfg.get("italic", False),
            cfg.get("colour", BLACK),
            cfg.get("hyperlink"),
        )
    return box


def chrome(slide, eyebrow, right_label, footer_left, footer_right):
    """Fixed chrome per Section 6.0A: left stripe, black header band, footer line."""
    # Left accent stripe, full height, amber
    add_rect(slide, Inches(0), Inches(0), Inches(0.1), H, AMBER_B)
    # Header band
    add_rect(slide, Inches(0), Inches(0), W, Inches(1.0), BLACK)
    add_text(slide, eyebrow, Inches(0.4), Inches(0.32), Inches(7.5), Inches(0.4),
             font_name=BODY_FONT, size=12, bold=True, colour=OFF_WHITE,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, right_label, Inches(5.4), Inches(0.32), Inches(7.5), Inches(0.4),
             font_name=BODY_FONT, size=12, bold=True, colour=TEAL,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # Footer line
    add_rect(slide, Inches(0.4), FOOTER_Y, Inches(12.53), Pt(0.75), BLACK)
    add_text(slide, footer_left, Inches(0.4), Inches(7.13), Inches(8), Inches(0.3),
             font_name=BODY_FONT, size=9, colour=BLACK, align=PP_ALIGN.LEFT)
    add_text(slide, footer_right, Inches(9.13), Inches(7.13), Inches(3.8), Inches(0.3),
             font_name=BODY_FONT, size=9, colour=BLACK, align=PP_ALIGN.RIGHT)


def stat_card_row(slide, cards, top, height=Inches(1.6), left_margin=Inches(0.4), right_margin=Inches(0.4), gap=Inches(0.15)):
    """cards: list of dicts {number, label, source}. Evenly spaced black tiles with teal top edge."""
    n = len(cards)
    total_w = W - left_margin - right_margin
    card_w = (total_w - gap * (n - 1)) / n
    x = left_margin
    for c in cards:
        add_rect(slide, x, top, card_w, height, BLACK)
        add_rect(slide, x, top, card_w, Inches(0.06), TEAL)
        add_text(slide, c["number"], x + Inches(0.12), top + Inches(0.16), card_w - Inches(0.24), Inches(0.5),
                  font_name=HEADING_FONT, size=28, bold=True, colour=WHITE, align=PP_ALIGN.LEFT)
        add_multiline(slide, [
            {"text": c["label"], "size": 10.5, "colour": OFF_WHITE, "bold": False, "line_spacing": 1.0},
        ], x + Inches(0.12), top + Inches(0.72), card_w - Inches(0.24), Inches(0.55))
        add_text(slide, c["source"], x + Inches(0.12), top + height - Inches(0.28), card_w - Inches(0.24), Inches(0.24),
                  font_name=BODY_FONT, size=7.5, colour=TEAL, align=PP_ALIGN.LEFT)
        x += card_w + gap

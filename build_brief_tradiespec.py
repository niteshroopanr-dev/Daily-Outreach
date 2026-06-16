"""
ProfitPulse Brief Builder
Target: TradieSpec | Date: 17 Jun 2026
Three-slide prospect-facing deck. Brand colours only. Zero dashes.
White background house style per Section 6 spec.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

# Brand colours
BLACK     = RGBColor(0x00, 0x00, 0x00)
TEAL      = RGBColor(0x01, 0xA2, 0x96)
AMBER_B   = RGBColor(0xF8, 0xC8, 0x06)
AMBER_D   = RGBColor(0xF6, 0xA1, 0x02)
GOLD      = RGBColor(0xE3, 0xA7, 0x12)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xE6, 0xE5, 0xDE)
MID_GREY  = RGBColor(0x66, 0x66, 0x66)
DRK_GREY  = RGBColor(0x33, 0x33, 0x33)
MUTED     = RGBColor(0x99, 0x99, 0x99)
TEAL_MID  = RGBColor(0x01, 0x6E, 0x65)

W = Inches(13.333)
H = Inches(7.5)

DATE    = "17 Jun 2026"
COMPANY = "TradieSpec"

QUEST_CLEAN = "profit-pulse.com.au/full-suite-of-products"
QUEST_FULL  = "https://profit-pulse.com.au/full-suite-of-products"
STRIPE_URL  = "https://buy.stripe.com/5kQ9AU7OU9io95m8sn3ks0l"
BOOKING_URL = ("https://bookings.cloud.microsoft/book/"
               "ProfitPulse1@profit-pulse.com.au/?ismsaljsauthenabled=true")


# ── Helpers ───────────────────────────────────────────────────────────────────

def bg_white(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def rect(slide, left, top, width, height, fill_colour, line_colour=None, lw=None):
    shp = slide.shapes.add_shape(1, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_colour
    if line_colour:
        shp.line.color.rgb = line_colour
        if lw:
            shp.line.width = lw
    else:
        shp.line.fill.background()
    return shp


def txt(slide, text, left, top, width, height,
        font="Arial", size=12, bold=False, colour=BLACK,
        align=PP_ALIGN.LEFT, italic=False, wrap=True):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = colour
    return tb


def multiline(slide, lines, left, top, width, height,
              font="Arial", default_size=12, default_colour=BLACK,
              default_bold=False, align=PP_ALIGN.LEFT, spacing_after=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        if isinstance(line, str):
            cfg = {"text": line, "size": default_size,
                   "colour": default_colour, "bold": default_bold,
                   "italic": False, "font": font}
        else:
            cfg = {
                "text":   line.get("text", ""),
                "size":   line.get("size", default_size),
                "colour": line.get("colour", default_colour),
                "bold":   line.get("bold", default_bold),
                "italic": line.get("italic", False),
                "font":   line.get("font", font),
            }
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if spacing_after:
            p.space_after = Pt(spacing_after)
        r = p.add_run()
        r.text = cfg["text"]
        r.font.name = cfg["font"]
        r.font.size = Pt(cfg["size"])
        r.font.bold = cfg["bold"]
        r.font.italic = cfg["italic"]
        r.font.color.rgb = cfg["colour"]
    return tb


def add_hyperlink(slide, tb, paragraph_idx, text, url,
                  font="Arial", size=12, colour=TEAL, bold=True):
    """Replace the text in an existing paragraph with a hyperlinked run."""
    tf = tb.text_frame
    p = tf.paragraphs[paragraph_idx]
    # Clear existing runs
    for r in list(p.runs):
        p._p.remove(r._r)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    rId = slide.part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True
    )
    rPr = run._r.get_or_add_rPr()
    hlinkClick = etree.SubElement(rPr, qn('a:hlinkClick'))
    hlinkClick.set(qn('r:id'), rId)
    return run


def header_band(slide, eyebrow, company):
    """Black header band with eyebrow left, company right."""
    rect(slide, Inches(0.1), Inches(0), W - Inches(0.1), Inches(0.88), BLACK)
    txt(slide, eyebrow,
        left=Inches(0.25), top=Inches(0.2), width=Inches(8), height=Inches(0.5),
        font="Arial", size=11, bold=True, colour=OFF_WHITE)
    txt(slide, company,
        left=Inches(9.5), top=Inches(0.2), width=Inches(3.6), height=Inches(0.5),
        font="Arial", size=11, bold=True, colour=OFF_WHITE,
        align=PP_ALIGN.RIGHT)


def amber_stripe(slide):
    """Thin amber vertical stripe on the far left."""
    rect(slide, Inches(0), Inches(0), Inches(0.1), H, AMBER_D)


def footer(slide):
    """Thin footer with prepared-by line and date."""
    rect(slide, Inches(0.1), Inches(7.18), W - Inches(0.1), Pt(1), MUTED)
    txt(slide,
        "Prepared by Nitesh Roopa CA, Managing Partner and Founder, "
        "ProfitPulse, Profit-Pulse.com.au",
        left=Inches(0.2), top=Inches(7.22), width=Inches(9), height=Inches(0.28),
        font="Arial", size=8, colour=MID_GREY)
    txt(slide, DATE,
        left=Inches(10.5), top=Inches(7.22), width=Inches(2.6), height=Inches(0.28),
        font="Arial", size=8, colour=MID_GREY, align=PP_ALIGN.RIGHT)


def stat_card(slide, left, top, width, height, number, label_line1, label_line2, source):
    """Black tile with teal top accent, large number, label, source."""
    # Card background
    rect(slide, left, top, width, height, BLACK)
    # Teal top accent (thin stripe)
    rect(slide, left, top, width, Pt(4), TEAL)
    # Number
    txt(slide, number,
        left=left + Inches(0.12), top=top + Inches(0.15),
        width=width - Inches(0.24), height=Inches(0.6),
        font="Georgia", size=26, bold=True, colour=WHITE)
    # Label line 1
    txt(slide, label_line1,
        left=left + Inches(0.12), top=top + Inches(0.72),
        width=width - Inches(0.24), height=Inches(0.28),
        font="Arial", size=10, bold=False, colour=OFF_WHITE)
    # Label line 2
    if label_line2:
        txt(slide, label_line2,
            left=left + Inches(0.12), top=top + Inches(0.98),
            width=width - Inches(0.24), height=Inches(0.25),
            font="Arial", size=9, bold=False, colour=OFF_WHITE)
    # Source line
    src_top = top + Inches(1.22) if label_line2 else top + Inches(0.98)
    txt(slide, source,
        left=left + Inches(0.12), top=src_top,
        width=width - Inches(0.24), height=Inches(0.25),
        font="Arial", size=7, bold=False, colour=TEAL_MID, italic=True)


# ═════════════════════════════════════════════════════════════════════════════
# PRESENTATION
# ═════════════════════════════════════════════════════════════════════════════
prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1: COMMERCIAL INTELLIGENCE BRIEF
# ─────────────────────────────────────────────────────────────────────────────
s1 = prs.slides.add_slide(blank)
bg_white(s1)
amber_stripe(s1)
header_band(s1, "COMMERCIAL INTELLIGENCE BRIEF", "PROFITPULSE")

# Company name (large, serif)
txt(s1, "TradieSpec",
    left=Inches(0.2), top=Inches(0.95), width=Inches(9), height=Inches(0.85),
    font="Georgia", size=44, bold=True, colour=BLACK)

# Descriptor
txt(s1,
    "Rent to own and flexible trade vehicle hire  |  Peakhurst, Sydney NSW",
    left=Inches(0.2), top=Inches(1.78), width=Inches(10), height=Inches(0.35),
    font="Arial", size=12, colour=TEAL)

# ── Stat cards ────────────────────────────────────────────────────────────────
CARD_TOP = Inches(2.2)
CARD_H   = Inches(1.62)
CARD_W   = Inches(2.12)
GAP      = Inches(0.065)
START_X  = Inches(0.2)

cards = [
    ("$15.3M",  "Revenue",             "",                      "Smart50 2025, SmartCompany"),
    ("57%",     "Three year average",  "revenue growth",        "Smart50 2025, SmartCompany"),
    ("#18",     "Smart50 2025",        "rank of 50",            "SmartCompany, Nov 2025"),
    ("22",      "Team members",        "",                      "Smart50 2025, SmartCompany"),
    ("600+",    "Trade vehicles",      "in fleet",              "Company public statements"),
    ("2,500+",  "Trade businesses",    "served since 2018",     "Company public statements"),
]

for i, (num, l1, l2, src) in enumerate(cards):
    stat_card(s1,
              left=START_X + i * (CARD_W + GAP),
              top=CARD_TOP,
              width=CARD_W,
              height=CARD_H,
              number=num, label_line1=l1, label_line2=l2, source=src)

# ── Key Commercial Signals ────────────────────────────────────────────────────
SIG_TOP = Inches(3.95)

txt(s1, "KEY COMMERCIAL SIGNALS",
    left=Inches(0.2), top=SIG_TOP, width=Inches(5), height=Inches(0.28),
    font="Arial", size=10, bold=True, colour=TEAL)

signals = [
    ("$13 million raised to fund national fleet expansion, with stated goal to double the fleet within 12 months",
     "Business Daily Media, 2023"),
    ("National operations now span Sydney, Melbourne, Brisbane, Perth, Central Coast and Newcastle",
     "tradiespec.com.au"),
    ("CEO Tim Cullen publicly stated ambition to make TradieSpec Australia's largest trade vehicle provider",
     "Business Daily Media, 2023"),
    ("57% three year average growth places TradieSpec among Australia's top 20 fastest growing SMEs",
     "SmartCompany Smart50 2025"),
    ("Also ranked 9th in SmartCompany Smart50 in 2022, confirming a sustained high growth trajectory",
     "SmartCompany Smart50 2022"),
    ("Founded November 2018 and now serves over 2,500 trade businesses across four Australian states",
     "Company public statements"),
]

sig_y = SIG_TOP + Inches(0.32)
for sig_text, sig_source in signals:
    # Bullet dot
    rect(s1, Inches(0.2), sig_y + Inches(0.09), Inches(0.07), Inches(0.07), TEAL)
    # Signal text
    txt(s1, sig_text,
        left=Inches(0.35), top=sig_y, width=Inches(7.7), height=Inches(0.27),
        font="Arial", size=10, colour=BLACK)
    # Source
    txt(s1, "Source: " + sig_source,
        left=Inches(8.2), top=sig_y + Inches(0.04), width=Inches(4.9), height=Inches(0.24),
        font="Arial", size=8, colour=MUTED, italic=True)
    sig_y += Inches(0.48)

footer(s1)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2: THE OPPORTUNITY — THREE COMMERCIAL OBSERVATIONS
# ─────────────────────────────────────────────────────────────────────────────
s2 = prs.slides.add_slide(blank)
bg_white(s2)
amber_stripe(s2)
header_band(s2, "THE OPPORTUNITY", COMPANY)

# Subtitle
txt(s2,
    "Three commercial observations from ProfitPulse",
    left=Inches(0.2), top=Inches(0.93), width=Inches(10), height=Inches(0.35),
    font="Arial", size=12, colour=MID_GREY, italic=True)

# ── Three columns ─────────────────────────────────────────────────────────────
COL_TOP = Inches(1.35)
COL_H   = Inches(5.3)
COL_W   = Inches(4.25)
COL_GAP = Inches(0.13)
COL_X   = [Inches(0.13), Inches(0.13) + COL_W + COL_GAP,
           Inches(0.13) + 2 * (COL_W + COL_GAP)]
COL_FILLS = [TEAL, BLACK, GOLD]
NUM_COL   = [WHITE, WHITE, BLACK]
HDR_COL   = [WHITE, AMBER_B, BLACK]
BODY_COL  = [WHITE, OFF_WHITE, BLACK]

observations = [
    {
        "idx": "01",
        "header": "Capital allocated at scale without a return map",
        "body": (
            "TradieSpec has raised $13 million and deployed it across a fleet of over 600 vehicles "
            "in four capital cities. Revenue sits at $15.3 million with 22 people "
            "running operations across six locations.\n\n"
            "The central financial question is not how fast to grow but where each "
            "dollar of fleet capital generates the strongest return. Which markets lead "
            "on utilisation? Which vehicle categories carry the highest margin? Which "
            "customer segments convert from hire to ownership most often?\n\n"
            "Without a formal capital allocation map, the risk in the next expansion "
            "round is investing in the next city rather than the best city. A Capital "
            "Allocation Review produces exactly that map before the next commitment."
        ),
    },
    {
        "idx": "02",
        "header": "Working capital dynamics compound at growth pace",
        "body": (
            "The rent to own model creates a layered cash flow structure that grows "
            "harder to manage as scale increases. Fleet acquisition is a large upfront "
            "outflow. Rental income is recurring but spread across over 2,500 accounts "
            "at varying intervals. Ownership transitions shift the cash profile of each "
            "account at the point of conversion.\n\n"
            "With $13 million in fleet financing likely carrying covenants and the "
            "business growing at 57 percent, the gap between reported profit and "
            "available cash can widen quickly if the 13 week cash flow view is not "
            "actively managed.\n\n"
            "A Fractional CFO Partnership or a 13 Week Cash Flow Build gives the "
            "team a clear view of the cash cycle before the next fleet tranche is drawn."
        ),
    },
    {
        "idx": "03",
        "header": "Senior financial leadership absent at a decisive stage",
        "body": (
            "At $15.3 million in revenue, $13 million in fleet financing, and an "
            "ambition to become Australia's largest trade vehicle provider, TradieSpec "
            "is making decisions that carry material financial consequence. Market entry "
            "sequencing, covenant management, fleet doubling, and a potential next "
            "capital raise all require senior financial thinking.\n\n"
            "With 22 employees, a full time CFO hire is premature. But leaving those "
            "decisions to the founding team without a senior finance voice at the table "
            "carries its own risk.\n\n"
            "The Fractional CFO Partnership places a senior financial partner at the "
            "table on a monthly cadence, at exactly the stage where it matters most."
        ),
    },
]

for col_i, obs in enumerate(observations):
    cx = COL_X[col_i]
    fill = COL_FILLS[col_i]
    rect(s2, cx, COL_TOP, COL_W, COL_H, fill)

    # Index number
    txt(s2, obs["idx"],
        left=cx + Inches(0.18), top=COL_TOP + Inches(0.18),
        width=COL_W - Inches(0.3), height=Inches(0.55),
        font="Georgia", size=30, bold=True, colour=NUM_COL[col_i])

    # Observation header
    txt(s2, obs["header"],
        left=cx + Inches(0.18), top=COL_TOP + Inches(0.78),
        width=COL_W - Inches(0.3), height=Inches(0.65),
        font="Georgia", size=13, bold=True, colour=HDR_COL[col_i], wrap=True)

    # Divider
    rect(s2,
         cx + Inches(0.18),
         COL_TOP + Inches(1.46),
         COL_W - Inches(0.36),
         Pt(1),
         AMBER_D if col_i == 1 else (WHITE if col_i == 0 else BLACK))

    # Body text
    multiline(s2,
              [{"text": obs["body"], "size": 10,
                "colour": BODY_COL[col_i], "bold": False}],
              left=cx + Inches(0.18), top=COL_TOP + Inches(1.58),
              width=COL_W - Inches(0.3), height=Inches(3.55),
              default_size=10, default_colour=BODY_COL[col_i])

# Warm closing line
txt(s2,
    "These are observations offered in good faith. TradieSpec has built something "
    "genuinely impressive. The question is simply whether the financial architecture "
    "matches the ambition.",
    left=Inches(0.2), top=Inches(6.73), width=Inches(13), height=Inches(0.38),
    font="Arial", size=10, colour=MID_GREY, italic=True)

footer(s2)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3: THE RECOMMENDATION AND HOW TO START
# ─────────────────────────────────────────────────────────────────────────────
s3 = prs.slides.add_slide(blank)
bg_white(s3)
amber_stripe(s3)
header_band(s3, "THE RECOMMENDATION", COMPANY)

# ── LEFT COLUMN: Recommendation and CTAs ──────────────────────────────────────
LX = Inches(0.2)
LW = Inches(7.9)

# Service name
txt(s3, "Capital Allocation Review",
    left=LX, top=Inches(0.97), width=LW, height=Inches(0.58),
    font="Georgia", size=22, bold=True, colour=BLACK)

# Price
txt(s3, "$7,500 one off",
    left=LX, top=Inches(1.52), width=LW, height=Inches(0.35),
    font="Arial", size=14, bold=True, colour=TEAL)

# What it does
txt(s3,
    "An independent review of where capital is deployed across the vehicle fleet, "
    "markets, and customer segments, against the return each generates. Produces "
    "a prioritised redeployment plan with expected ROI on each move, giving the "
    "team a clear investment thesis before committing the next round of fleet capital.",
    left=LX, top=Inches(1.93), width=LW, height=Inches(0.9),
    font="Arial", size=11, colour=DRK_GREY, wrap=True)

# Divider line
rect(s3, LX, Inches(2.9), LW, Pt(1.5), TEAL)

# Step one block
rect(s3, LX, Inches(2.98), LW, Inches(1.55),
     RGBColor(0xF2, 0xFB, 0xFA))
rect(s3, LX, Inches(2.98), Pt(4), Inches(1.55), TEAL)

txt(s3, "Step one: answer a few quick questions",
    left=LX + Inches(0.15), top=Inches(3.04), width=LW - Inches(0.3), height=Inches(0.35),
    font="Arial", size=12, bold=True, colour=BLACK)
txt(s3, "See the solutions matched to your size and industry.",
    left=LX + Inches(0.15), top=Inches(3.4), width=LW - Inches(0.3), height=Inches(0.28),
    font="Arial", size=11, colour=DRK_GREY)

# Questionnaire URL (clean, no UTM on brief per Section 6.4)
q_tb = txt(s3, QUEST_CLEAN,
           left=LX + Inches(0.15), top=Inches(3.7), width=LW - Inches(0.3), height=Inches(0.3),
           font="Arial", size=11, colour=TEAL, bold=False)
# Make it a hyperlink to the clean URL
add_hyperlink(s3, q_tb, 0, QUEST_CLEAN, QUEST_FULL,
              font="Arial", size=11, colour=TEAL, bold=False)

# Divider
rect(s3, LX, Inches(4.6), LW, Pt(1.5), AMBER_D)

# Direct purchase CTA (Stripe link hidden behind words)
cta_tb = slide_shapes = txt(s3, "",
                              left=LX, top=Inches(4.67), width=LW, height=Inches(0.5))
# Add the CTA text as a hyperlinked run directly
tf = cta_tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run = p.add_run()
run.text = "Purchase the suggested product now to get started"
run.font.name = "Arial"
run.font.size = Pt(12)
run.font.bold = True
run.font.color.rgb = AMBER_D
rId = s3.part.relate_to(
    STRIPE_URL,
    'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
    is_external=True
)
rPr = run._r.get_or_add_rPr()
hl = etree.SubElement(rPr, qn('a:hlinkClick'))
hl.set(qn('r:id'), rId)

# Already know this is the service caption
txt(s3, "Capital Allocation Review  |  $7,500 one off  |  ProfitPulse verified price",
    left=LX, top=Inches(5.18), width=LW, height=Inches(0.28),
    font="Arial", size=9, colour=MUTED, italic=True)

# Discovery call
txt(s3, "Prefer a conversation first?",
    left=LX, top=Inches(5.55), width=LW, height=Inches(0.28),
    font="Arial", size=11, colour=DRK_GREY)

bk_tb = txt(s3, "Book a complimentary discovery call",
            left=LX, top=Inches(5.84), width=LW, height=Inches(0.3),
            font="Arial", size=11, colour=TEAL, bold=True)
add_hyperlink(s3, bk_tb, 0, "Book a complimentary discovery call",
              BOOKING_URL, font="Arial", size=11, colour=TEAL, bold=True)

# Supporting services
txt(s3,
    "Supporting services identified: B1 Fractional CFO Partnership  |  F1 Capital Raise Feasibility",
    left=LX, top=Inches(6.28), width=LW, height=Inches(0.28),
    font="Arial", size=8, colour=MUTED, italic=True)

# ── RIGHT COLUMN: About Nitesh ────────────────────────────────────────────────
RX = Inches(8.35)
RW = Inches(4.8)

rect(s3, RX, Inches(0.95), RW, Inches(6.15),
     RGBColor(0xF7, 0xF7, 0xF5))
rect(s3, RX, Inches(0.95), Pt(4), Inches(6.15), TEAL)

txt(s3, "Nitesh Roopa",
    left=RX + Inches(0.2), top=Inches(1.05), width=RW - Inches(0.3), height=Inches(0.55),
    font="Georgia", size=20, bold=True, colour=BLACK)

txt(s3, "CA, Managing Partner  |  ProfitPulse",
    left=RX + Inches(0.2), top=Inches(1.58), width=RW - Inches(0.3), height=Inches(0.32),
    font="Arial", size=11, colour=TEAL)

rect(s3, RX + Inches(0.2), Inches(1.96), RW - Inches(0.4), Pt(1.5), TEAL)

creds = [
    "16 years across 4 countries",
    "52 deals executed and managed",
    "Largest single deal: USD 1.3 billion,",
    "  Cahora Bassa, Mozambique Government",
    "Total GRBT value in Queensland:",
    "  over AUD 10 billion",
]
cy = Inches(2.08)
for line in creds:
    txt(s3, line,
        left=RX + Inches(0.2), top=cy, width=RW - Inches(0.3), height=Inches(0.3),
        font="Arial", size=10, colour=DRK_GREY)
    cy += Inches(0.3)

rect(s3, RX + Inches(0.2), cy + Inches(0.06), RW - Inches(0.4), Pt(1.5), AMBER_D)
cy += Inches(0.22)

contact = [
    ("Profit-Pulse.com.au",           BLACK),
    ("Nitesh@Profit-Pulse.com.au",    TEAL),
    ("+61 411 876 267",               BLACK),
    ("linkedin.com/in/nitesh-roopa-77594163", MID_GREY),
]
for ctxt, ccol in contact:
    txt(s3, ctxt,
        left=RX + Inches(0.2), top=cy, width=RW - Inches(0.3), height=Inches(0.3),
        font="Arial", size=10, colour=ccol)
    cy += Inches(0.3)

footer(s3)


# ═════════════════════════════════════════════════════════════════════════════
out_path = "/home/user/Daily-Outreach/Out-reach efforts/Brief_TradieSpec_17Jun2026.pptx"
prs.save(out_path)
print(f"PPTX saved: {out_path}")
